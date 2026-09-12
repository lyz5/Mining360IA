from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


class BusinessReviewExportService:
    """Build a versioned executive workbook from already-authorized payloads."""

    @staticmethod
    def build(*, context, confidence, metrics, portfolio):
        workbook = Workbook()
        summary = workbook.active
        summary.title = "Executive Overview"
        portfolio_sheet = workbook.create_sheet("Portfolio")
        metadata = workbook.create_sheet("Review Context")

        summary.append(["Metric", "Value"])
        for code, label in (
            ("mining_revenue_ytd_eur", "Mining Revenue YTD (EUR)"),
            ("mining_revenue_previous_year_eur", "Previous Year Revenue (EUR)"),
            ("fleet_count", "Fleet Count"),
            ("revenue_per_equipment_eur", "Revenue per Equipment (EUR)"),
            ("revenue_assigned_eur", "Revenue Assigned (EUR)"),
            ("unallocated_revenue_eur", "Unallocated Revenue (EUR)"),
            ("revenue_coverage_pct", "Revenue Coverage (%)"),
            ("fleet_coverage_pct", "Fleet Coverage (%)"),
            ("account_coverage_pct", "Account Coverage (%)"),
        ):
            summary.append([label, metrics.get(code) if metrics.get(code) is not None else "Not available"])
        summary.append([])
        summary.append(["Revenue Lens", "Value (EUR)"])
        for code, value in (metrics.get("revenue_by_lob") or {}).items():
            summary.append([code, value])

        portfolio_sheet.append(["MineSite", "Classification", "Fleet", "Revenue (EUR)", "Revenue per Equipment (EUR)", "Account Count"])
        for item in portfolio:
            portfolio_sheet.append([item.get("name"), item.get("classification") or "Not classified", item.get("fleet"), item.get("revenue"), item.get("revenue_per_equipment") or "Not available", item.get("account_count")])

        metadata.append(["Field", "Value"])
        for key, value in context.items():
            metadata.append([key, str(value) if value is not None else "Not available"])
        metadata.append(["data_confidence", confidence.get("status", "Not available")])
        metadata.append(["confidence_rule_version", confidence.get("rule_version", "Not available")])
        for warning in confidence.get("warnings", []):
            metadata.append(["warning", warning])

        for sheet in workbook.worksheets:
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="17223B")
            sheet.freeze_panes = "A2"
            for column in sheet.columns:
                width = min(55, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
                sheet.column_dimensions[column[0].column_letter].width = width

        stream = BytesIO()
        workbook.save(stream)
        return stream.getvalue()
