from __future__ import annotations

import csv
import hashlib
import json
from collections import defaultdict
from pathlib import Path

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from .models import (
    EquipmentModelReference,
    EquipmentPrefixModelReference,
    EquipmentProductGroupReference,
    EquipmentReferenceImportRun,
    EquipmentSerialReference,
)


SERIAL_FIELDS = (
    "serial_prefix", "make", "model", "product_family", "dealer_customer_name",
    "ownership_status", "model_year", "subscription_status", "source_file_name",
    "source_row_number", "source_hash", "active", "import_run", "source_last_seen_at",
)
PREFIX_FIELDS = (
    "brand", "prefix", "model", "parent_product_family", "product_family",
    "equipment_type", "source_created_by", "ambiguous_prefix", "source_file_name",
    "source_row_number", "source_hash", "active", "import_run", "source_last_seen_at",
)
PRODUCT_GROUP_FIELDS = (
    "description", "priority", "source_created_by", "source_file_name",
    "source_row_number", "source_hash", "active", "import_run", "source_last_seen_at",
)
MODEL_FIELDS = (
    "model", "normalized_model", "brand", "family", "priority", "equipment_type",
    "description", "source_status", "source_created_by", "product_group",
    "source_file_name", "source_row_number", "source_hash", "active", "import_run",
    "source_last_seen_at",
)


def _text(value):
    return str(value or "").strip()


def _code(value):
    return "".join(character for character in _text(value).upper() if character.isalnum())


def _integer(value):
    try:
        return max(0, int(float(value or 0)))
    except (TypeError, ValueError):
        return 0


def _brand(value):
    clean = _text(value)
    normalized = clean.casefold()
    if normalized in {"cat", "caterpillar", "caterpillar inc"}:
        return "CAT"
    if normalized in {"epi", "epr", "epiroc"}:
        return "Epiroc"
    return clean


def _payload_hash(payload):
    return hashlib.sha256(json.dumps(payload, sort_keys=True, default=str).encode("utf-8")).hexdigest()


def _file_hash(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _deactivate_missing(model, current_keys, key_function):
    missing_ids = [item.pk for item in model.objects.filter(active=True) if key_function(item) not in current_keys]
    count = 0
    for offset in range(0, len(missing_ids), 500):
        count += model.objects.filter(pk__in=missing_ids[offset:offset + 500]).update(active=False)
    return count


class EquipmentReferenceImportService:
    @staticmethod
    def _read_serials(path):
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = [_text(value) for value in next(rows)]
        required = {"Asset Serial Number", "Make", "Model", "Product Family"}
        if not required.issubset(headers):
            raise ValueError(f"Asset List is missing required columns: {', '.join(sorted(required - set(headers)))}")
        payloads = {}
        for row_number, values in enumerate(rows, start=2):
            source = dict(zip(headers, values))
            serial_number = _code(source.get("Asset Serial Number"))
            if not serial_number:
                continue
            payload = {
                "serial_number": serial_number,
                "serial_prefix": serial_number[:3],
                "make": _brand(source.get("Make")),
                "model": _text(source.get("Model")),
                "product_family": _text(source.get("Product Family")),
                "dealer_customer_name": _text(source.get("Dealer Customer Name")),
                "ownership_status": _text(source.get("Ownership Status")),
                "model_year": _text(source.get("Model Year")),
                "subscription_status": _text(source.get("Subscription Status")),
                "source_file_name": Path(path).name,
                "source_row_number": row_number,
                "active": True,
            }
            payload["source_hash"] = _payload_hash(payload)
            payloads[serial_number] = payload
        workbook.close()
        return payloads

    @staticmethod
    def _read_prefixes(path):
        raw_payloads = {}
        models_by_prefix = defaultdict(set)
        with Path(path).open("r", encoding="utf-8-sig", newline="") as source_file:
            reader = csv.DictReader(source_file)
            required = {"Brand", "Parent Product Family", "Product Family", "Models", "Prefixes"}
            missing = required - set(reader.fieldnames or [])
            if missing:
                raise ValueError(f"Prefix file is missing required columns: {', '.join(sorted(missing))}")
            for row_number, source in enumerate(reader, start=2):
                brand = _brand(source.get("Brand"))
                prefix = _code(source.get("Prefixes"))
                model = _text(source.get("Models"))
                if not prefix or not model:
                    continue
                key = (
                    brand,
                    prefix,
                    model,
                    _text(source.get("Parent Product Family")),
                    _text(source.get("Product Family")),
                )
                models_by_prefix[(brand, prefix)].add(model.casefold())
                raw_payloads[key] = {
                    "brand": brand,
                    "prefix": prefix,
                    "model": model,
                    "parent_product_family": key[3],
                    "product_family": key[4],
                    "equipment_type": _text(source.get("EquipType")),
                    "source_created_by": _text(source.get("Created by")),
                    "source_file_name": Path(path).name,
                    "source_row_number": row_number,
                    "active": True,
                }
        for key, payload in raw_payloads.items():
            payload["ambiguous_prefix"] = len(models_by_prefix[(payload["brand"], payload["prefix"])]) > 1
            payload["source_hash"] = _payload_hash(payload)
        return raw_payloads

    @staticmethod
    def _read_product_groups(path):
        payloads = {}
        with Path(path).open("r", encoding="utf-8-sig", newline="") as source_file:
            reader = csv.DictReader(source_file)
            required = {"Priority", "Product Group Code", "Product Group Description"}
            missing = required - set(reader.fieldnames or [])
            if missing:
                raise ValueError(f"Product Group file is missing required columns: {', '.join(sorted(missing))}")
            for row_number, source in enumerate(reader, start=2):
                code = _code(source.get("Product Group Code"))
                if not code:
                    continue
                payload = {
                    "code": code,
                    "description": _text(source.get("Product Group Description")),
                    "priority": _integer(source.get("Priority")),
                    "source_created_by": _text(source.get("Created by")),
                    "source_file_name": Path(path).name,
                    "source_row_number": row_number,
                    "active": True,
                }
                payload["source_hash"] = _payload_hash(payload)
                payloads[code] = payload
        return payloads

    @staticmethod
    def _read_models(path):
        payloads = {}
        with Path(path).open("r", encoding="utf-8-sig", newline="") as source_file:
            reader = csv.DictReader(source_file)
            required = {"ID", "Model", "Family", "Brand", "Prime Movers"}
            missing = required - set(reader.fieldnames or [])
            if missing:
                raise ValueError(f"Equipment Models file is missing required columns: {', '.join(sorted(missing))}")
            for row_number, source in enumerate(reader, start=2):
                source_record_id = _text(source.get("ID"))
                model = _text(source.get("Model"))
                if not source_record_id or not model:
                    continue
                payload = {
                    "source_record_id": source_record_id,
                    "model": model,
                    "normalized_model": _code(model),
                    "brand": _brand(source.get("Brand")),
                    "family": _text(source.get("Family")),
                    "priority": _integer(source.get("Priority")),
                    "equipment_type": _text(source.get("Type")),
                    "description": _text(source.get("Description")),
                    "source_status": _text(source.get("Status")),
                    "source_created_by": _text(source.get("Created By")),
                    "product_group_code": _code(source.get("Prime Movers")),
                    "source_file_name": Path(path).name,
                    "source_row_number": row_number,
                    "active": True,
                }
                payload["source_hash"] = _payload_hash(payload)
                payloads[source_record_id] = payload
        return payloads

    @classmethod
    def import_files(
        cls, *, asset_list_path, prefix_csv_path,
        product_group_csv_path=None, model_csv_path=None, user=None,
    ):
        asset_path = Path(asset_list_path)
        prefix_path = Path(prefix_csv_path)
        product_group_path = Path(product_group_csv_path) if product_group_csv_path else None
        model_path = Path(model_csv_path) if model_csv_path else None
        if not asset_path.is_file():
            raise FileNotFoundError(f"Asset List not found: {asset_path}")
        if not prefix_path.is_file():
            raise FileNotFoundError(f"Prefix file not found: {prefix_path}")
        if product_group_path and not product_group_path.is_file():
            raise FileNotFoundError(f"Product Group file not found: {product_group_path}")
        if model_path and not model_path.is_file():
            raise FileNotFoundError(f"Equipment Models file not found: {model_path}")

        run = EquipmentReferenceImportRun.objects.create(
            asset_source_name=asset_path.name,
            prefix_source_name=prefix_path.name,
            product_group_source_name=product_group_path.name if product_group_path else "",
            model_source_name=model_path.name if model_path else "",
            imported_by=user,
        )
        try:
            serial_payloads = cls._read_serials(asset_path)
            prefix_payloads = cls._read_prefixes(prefix_path)
            product_group_payloads = cls._read_product_groups(product_group_path) if product_group_path else None
            model_payloads = cls._read_models(model_path) if model_path else None
            now = timezone.now()
            created = updated = unchanged = deactivated = 0
            with transaction.atomic():
                serial_existing = {
                    item.serial_number: item
                    for item in EquipmentSerialReference.objects.all()
                }
                serial_creates = []
                serial_updates = []
                for key, payload in serial_payloads.items():
                    item = serial_existing.get(key)
                    if item is None:
                        serial_creates.append(EquipmentSerialReference(
                            import_run=run, source_last_seen_at=now, **payload
                        ))
                        created += 1
                    elif item.source_hash != payload["source_hash"] or not item.active:
                        for field, value in payload.items():
                            setattr(item, field, value)
                        item.import_run = run
                        item.source_last_seen_at = now
                        serial_updates.append(item)
                        updated += 1
                    else:
                        item.import_run = run
                        item.source_last_seen_at = now
                        serial_updates.append(item)
                        unchanged += 1
                EquipmentSerialReference.objects.bulk_create(serial_creates, batch_size=500)
                EquipmentSerialReference.objects.bulk_update(
                    serial_updates, SERIAL_FIELDS, batch_size=500
                )
                deactivated += _deactivate_missing(
                    EquipmentSerialReference, set(serial_payloads), lambda item: item.serial_number
                )

                prefix_existing = {
                    (item.brand, item.prefix, item.model, item.parent_product_family, item.product_family): item
                    for item in EquipmentPrefixModelReference.objects.all()
                }
                prefix_creates = []
                prefix_updates = []
                for key, payload in prefix_payloads.items():
                    item = prefix_existing.get(key)
                    if item is None:
                        prefix_creates.append(EquipmentPrefixModelReference(
                            import_run=run, source_last_seen_at=now, **payload
                        ))
                        created += 1
                    elif item.source_hash != payload["source_hash"] or not item.active:
                        for field, value in payload.items():
                            setattr(item, field, value)
                        item.import_run = run
                        item.source_last_seen_at = now
                        prefix_updates.append(item)
                        updated += 1
                    else:
                        item.import_run = run
                        item.source_last_seen_at = now
                        prefix_updates.append(item)
                        unchanged += 1
                EquipmentPrefixModelReference.objects.bulk_create(prefix_creates, batch_size=500)
                EquipmentPrefixModelReference.objects.bulk_update(
                    prefix_updates, PREFIX_FIELDS, batch_size=500
                )
                deactivated += _deactivate_missing(
                    EquipmentPrefixModelReference,
                    set(prefix_payloads),
                    lambda item: (
                        item.brand, item.prefix, item.model,
                        item.parent_product_family, item.product_family,
                    ),
                )

                if product_group_payloads is not None:
                    group_existing = {
                        item.code: item for item in EquipmentProductGroupReference.objects.all()
                    }
                    group_creates = []
                    group_updates = []
                    for key, payload in product_group_payloads.items():
                        item = group_existing.get(key)
                        if item is None:
                            group_creates.append(EquipmentProductGroupReference(
                                import_run=run, source_last_seen_at=now, **payload
                            ))
                            created += 1
                        elif item.source_hash != payload["source_hash"] or not item.active:
                            for field, value in payload.items():
                                setattr(item, field, value)
                            item.import_run = run
                            item.source_last_seen_at = now
                            group_updates.append(item)
                            updated += 1
                        else:
                            item.import_run = run
                            item.source_last_seen_at = now
                            group_updates.append(item)
                            unchanged += 1
                    EquipmentProductGroupReference.objects.bulk_create(group_creates, batch_size=100)
                    EquipmentProductGroupReference.objects.bulk_update(
                        group_updates, PRODUCT_GROUP_FIELDS, batch_size=100
                    )
                    deactivated += _deactivate_missing(
                        EquipmentProductGroupReference, set(product_group_payloads), lambda item: item.code
                    )

                unknown_group_codes = set()
                if model_payloads is not None:
                    groups_by_code = {
                        item.code: item for item in EquipmentProductGroupReference.objects.filter(active=True)
                    }
                    model_existing = {
                        item.source_record_id: item for item in EquipmentModelReference.objects.all()
                    }
                    model_creates = []
                    model_updates = []
                    for key, payload in model_payloads.items():
                        values = {field: value for field, value in payload.items() if field != "product_group_code"}
                        group_code = payload["product_group_code"]
                        product_group = groups_by_code.get(group_code)
                        if group_code and product_group is None:
                            unknown_group_codes.add(group_code)
                        item = model_existing.get(key)
                        if item is None:
                            model_creates.append(EquipmentModelReference(
                                product_group=product_group, import_run=run,
                                source_last_seen_at=now, **values,
                            ))
                            created += 1
                        elif item.source_hash != payload["source_hash"] or not item.active or item.product_group_id != getattr(product_group, "pk", None):
                            for field, value in values.items():
                                setattr(item, field, value)
                            item.product_group = product_group
                            item.import_run = run
                            item.source_last_seen_at = now
                            model_updates.append(item)
                            updated += 1
                        else:
                            item.import_run = run
                            item.source_last_seen_at = now
                            model_updates.append(item)
                            unchanged += 1
                    EquipmentModelReference.objects.bulk_create(model_creates, batch_size=500)
                    EquipmentModelReference.objects.bulk_update(
                        model_updates, MODEL_FIELDS, batch_size=500
                    )
                    deactivated += _deactivate_missing(
                        EquipmentModelReference, set(model_payloads), lambda item: item.source_record_id
                    )

                ambiguous = len({
                    (payload["brand"], payload["prefix"])
                    for payload in prefix_payloads.values() if payload["ambiguous_prefix"]
                })
                warnings = []
                if ambiguous:
                    warnings.append(
                        f"{ambiguous} prefixes map to several models and are excluded from automatic fallback."
                    )
                if unknown_group_codes:
                    warnings.append(
                        "Unknown product-group codes in Equipment Models: "
                        + ", ".join(sorted(unknown_group_codes))
                    )
                run.asset_source_hash = _file_hash(asset_path)
                run.prefix_source_hash = _file_hash(prefix_path)
                run.product_group_source_hash = _file_hash(product_group_path) if product_group_path else ""
                run.model_source_hash = _file_hash(model_path) if model_path else ""
                run.serial_records_read = len(serial_payloads)
                run.prefix_records_read = len(prefix_payloads)
                run.product_group_records_read = len(product_group_payloads or {})
                run.model_records_read = len(model_payloads or {})
                run.records_created = created
                run.records_updated = updated
                run.records_unchanged = unchanged
                run.records_deactivated = deactivated
                run.warnings_json = warnings
                run.status = "Completed with Warnings" if warnings else "Completed"
                run.completed_at = timezone.now()
                run.save()
            return run
        except Exception as exc:
            run.status = "Failed"
            run.error_message = str(exc)[:2000]
            run.completed_at = timezone.now()
            run.save(update_fields=["status", "error_message", "completed_at"])
            raise
