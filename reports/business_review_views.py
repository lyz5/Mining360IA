from __future__ import annotations

import json
from io import BytesIO
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import login_required
from django.db.models import Count, Q, Sum
from django.core.serializers.json import DjangoJSONEncoder
from django.http import HttpResponseForbidden, JsonResponse
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone
from django.views.decorators.http import require_http_methods
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .ai_feature_rollout import feature_enabled
from .access_control import is_platform_admin
from .business_mapping_normalization_service import normalize_business_name
from .business_mapping_source_service import MINING_DIVISION, MINING_REVENUE_LABELS, MINING_REVENUE_LOBS
from .business_opportunity_service import BusinessOpportunityRuleEngine, BusinessRiskService
from .business_review_access_service import business_review_enabled, filter_published_rows, has_business_review_permission
from .business_review_confidence_service import BusinessReviewDataConfidenceService
from .business_review_export_service import BusinessReviewExportService
from .business_review_snapshot_service import BusinessReviewSnapshotService
from .business_command_center_service import BusinessCommandCenterInputError, BusinessCommandCenterService
from .machine_sales_service import MachineSalesDetailService
from .parts_sales_service import PartsSalesDetailService
from .models import (
    BusinessAccount,
    BusinessDecision,
    BusinessOpportunity,
    BusinessReviewAction,
    BusinessReviewSavedView,
    BusinessCommandCenterUserVisit,
    BusinessCommandCenterWatchlist,
    BusinessRisk,
    EquipmentFleetAnalysis,
    MappingAuditLog,
    MappingPublication,
    MineSite,
    RevenueSourceSnapshot,
)


def _command_center_allowed(user):
    return (
        business_review_enabled(user)
        and feature_enabled("ENABLE_BUSINESS_COMMAND_CENTER", user)
        and has_business_review_permission(user, "view_business_command_center")
    )


@login_required
def business_review_landing(request):
    if _command_center_allowed(request.user):
        return business_command_center(request)
    return business_review(request)


@login_required
def business_command_center(request):
    if not _command_center_allowed(request.user):
        return HttpResponseForbidden("You do not have access to Business Overview.")
    legacy_requested = request.GET.get("ui") == "legacy" and is_platform_admin(request.user)
    v2_enabled = feature_enabled("ENABLE_BUSINESS_COMMAND_CENTER_V2", request.user)
    template_name = "reports/business_command_center_v2.html" if v2_enabled and not legacy_requested else "reports/business_command_center_legacy.html"
    return render(request, template_name, {
        "active_section": "business-review",
        "can_export": has_business_review_permission(request.user, "export_business_command_center"),
        "can_create_action": has_business_review_permission(request.user, "create_business_review_action"),
        "v2_enabled": v2_enabled,
        "features": {
            "customers": feature_enabled("ENABLE_BUSINESS_COMMAND_CENTER_CUSTOMERS", request.user),
            "countries": feature_enabled("ENABLE_BUSINESS_COMMAND_CENTER_COUNTRIES", request.user),
            "key_accounts": feature_enabled("ENABLE_BUSINESS_COMMAND_CENTER_KEY_ACCOUNTS", request.user),
            "last_visit": feature_enabled("ENABLE_BUSINESS_COMMAND_CENTER_LAST_VISIT", request.user),
            "attention": feature_enabled("ENABLE_BUSINESS_COMMAND_CENTER_ATTENTION", request.user),
            "watchlist": feature_enabled("ENABLE_BUSINESS_COMMAND_CENTER_WATCHLIST", request.user),
            "presentation": feature_enabled("ENABLE_BUSINESS_COMMAND_CENTER_PRESENTATION_MODE", request.user),
        },
    })


@login_required
def command_center_bootstrap_api(request):
    if not _command_center_allowed(request.user):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    try:
        return JsonResponse(BusinessCommandCenterService(request.user, request.GET).bootstrap())
    except BusinessCommandCenterInputError as exc:
        return JsonResponse({"ready": False, "status": "INVALID_CONTEXT", "message": str(exc)}, status=400)
    except Exception:
        return JsonResponse({
            "ready": False,
            "status": "SOURCE_UNAVAILABLE",
            "message": "Revenue data is temporarily unavailable. Please retry the view.",
        }, status=503)


@login_required
def command_center_revenue_explorer_api(request):
    if not _command_center_allowed(request.user):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    try:
        return JsonResponse(BusinessCommandCenterService(request.user, request.GET).revenue_explorer())
    except BusinessCommandCenterInputError as exc:
        return JsonResponse({"ready": False, "message": str(exc)}, status=400)
    except Exception:
        return JsonResponse({"ready": False, "message": "Revenue Explorer is temporarily unavailable."}, status=503)


def _command_center_entity_search(request, entity_type):
    if not _command_center_allowed(request.user):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    try:
        return JsonResponse(BusinessCommandCenterService(request.user, request.GET).search_entities(entity_type))
    except BusinessCommandCenterInputError as exc:
        return JsonResponse({"results": [], "message": str(exc)}, status=400)
    except Exception:
        return JsonResponse({"results": [], "message": "Business search is temporarily unavailable."}, status=503)


@login_required
def command_center_customer_search_api(request):
    return _command_center_entity_search(request, "customers")


@login_required
def command_center_fleet_api(request):
    if not _command_center_allowed(request.user):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    try:
        return JsonResponse(BusinessCommandCenterService(request.user, request.GET).entity_fleet(
            request.GET.get("dimension"), request.GET.get("entity_id")))
    except BusinessCommandCenterInputError as exc:
        return JsonResponse({"message": str(exc)}, status=400)
    except Exception:
        return JsonResponse({"message": "Fleet is temporarily unavailable."}, status=503)


@login_required
def command_center_key_account_search_api(request):
    return _command_center_entity_search(request, "key_accounts")


def _machine_sales_params(request):
    _revenue, _publication, _published_rows, selected_rows, period, _line, _run = BusinessCommandCenterService(
        request.user, request.GET
    )._context()
    params = request.GET.copy()
    params["start_date"] = period["start_date"].isoformat()
    params["end_date"] = period["end_date"].isoformat()
    if any(request.GET.get(key) for key in ("customer_ids", "customer_group_ids", "country_ids", "key_account_ids")):
        params["customer_codes"] = ",".join(sorted({
            str(code) for row in selected_rows for code in row.get("source_account_codes", []) if code
        }))
    return params, period


@login_required
def command_center_machine_sales_api(request):
    if not _command_center_allowed(request.user):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    try:
        params, _period = _machine_sales_params(request)
        return JsonResponse(MachineSalesDetailService(request.user, params).result())
    except (BusinessCommandCenterInputError, ValueError) as exc:
        return JsonResponse({"ready": False, "message": str(exc)}, status=400)
    except Exception:
        return JsonResponse({
            "ready": False,
            "message": "Machine Sales details are temporarily unavailable. Revenue metrics remain available.",
        }, status=503)


@login_required
def command_center_parts_sales_api(request):
    if not _command_center_allowed(request.user):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    try:
        params, _period = _machine_sales_params(request)
        return JsonResponse(PartsSalesDetailService(request.user, params).result())
    except (BusinessCommandCenterInputError, ValueError) as exc:
        return JsonResponse({"ready": False, "message": str(exc)}, status=400)
    except Exception:
        return JsonResponse({
            "ready": False,
            "message": "Parts Sales classification is temporarily unavailable. Revenue metrics remain available.",
        }, status=503)


def _excel_value(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


@login_required
def command_center_machine_sales_export_api(request):
    if not _command_center_allowed(request.user) or not has_business_review_permission(request.user, "export_business_command_center"):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    try:
        params, period = _machine_sales_params(request)
        params["page_size"] = str(MachineSalesDetailService.MAX_PAGE_SIZE)
        params["page"] = "1"
        first = MachineSalesDetailService(request.user, params).result()
        rows = list(first["results"])
        for page in range(2, min(first["pagination"]["pages"], 334) + 1):
            params["page"] = str(page)
            rows.extend(MachineSalesDetailService(request.user, params).result()["results"])
    except (BusinessCommandCenterInputError, ValueError) as exc:
        return JsonResponse({"detail": str(exc)}, status=400)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Machines Sold"
    headers = ["Latest Invoice Date", "Customer", "Customer Code", "Model", "Family", "Source Family", "Brand", "Serial Number", "Equipment", "Invoices", "Condition", "Machine Sale EUR", "Other Charges & Adjustments EUR", "Net Invoiced Revenue EUR", "Source Entries"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="11152F")
    for item in rows:
        sheet.append([_excel_value(value) for value in (
            item["business_date"], item["customer_name"], item["customer_code"], item["model_name"],
            item["family_code"], item["equipment_family"], item["brand"], item["serial_number"],
            item["equipment_code"], ", ".join(item["invoice_numbers"]), ", ".join(item["conditions"]),
            item["machine_sale_eur"], item["other_charges_eur"], item["net_revenue_eur"], item["transaction_count"],
        )])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [18, 34, 18, 18, 12, 25, 14, 20, 16, 28, 18, 20, 28, 24, 14]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    detail = workbook.create_sheet("Invoice Entries")
    detail_headers = ["Date", "Customer", "Customer Code", "Model", "Serial Number", "Equipment", "Invoice", "Classification", "Source Category", "Amount EUR"]
    detail.append(detail_headers)
    for cell in detail[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="11152F")
    for item in rows:
        for entry in item["entries"]:
            detail.append([_excel_value(value) for value in (
                entry["business_date"], item["customer_name"], item["customer_code"], item["model_name"],
                item["serial_number"], item["equipment_code"], entry["invoice_number"],
                entry["classification"], entry["product_category"], entry["amount_eur"],
            )])
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = detail.dimensions
    metadata = workbook.create_sheet("Context")
    metadata.append(["Period", period["label"]])
    metadata.append(["Start Date", period["start_date"].isoformat()])
    metadata.append(["End Date", period["end_date"].isoformat()])
    metadata.append(["Family", request.GET.get("family") or "All"])
    metadata.append(["Brand", request.GET.get("brand") or "All"])
    metadata.append(["Generated At", timezone.now().isoformat()])
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="Mining360_Machines_Sold.xlsx"'
    return response


@login_required
def command_center_parts_sales_export_api(request):
    if not _command_center_allowed(request.user) or not has_business_review_permission(request.user, "export_business_command_center"):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    try:
        params, period = _machine_sales_params(request)
        params["page_size"] = str(PartsSalesDetailService.MAX_PAGE_SIZE)
        params["page"] = "1"
        first = PartsSalesDetailService(request.user, params).result()
        rows = list(first["results"])
        for page in range(2, first["pagination"]["pages"] + 1):
            params["page"] = str(page)
            rows.extend(PartsSalesDetailService(request.user, params).result()["results"])
    except (BusinessCommandCenterInputError, ValueError) as exc:
        return JsonResponse({"detail": str(exc)}, status=400)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Parts Classification"
    headers = ["Rank", "Brand Group", "Major Class", "Major Description", "Minor Class", "PPC", "Revenue EUR", "Invoice Lines", "Distinct Parts"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="11152F")
    for item in rows:
        sheet.append([
            item["rank"], item["brand_group"], item["major_class"], item["major_description"], item["minor_class"],
            item["ppc"], item["revenue_eur"], item["line_count"], item["part_count"],
        ])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    metadata = workbook.create_sheet("Context")
    metadata.append(["Period", period["label"]])
    metadata.append(["Start Date", period["start_date"].isoformat()])
    metadata.append(["End Date", period["end_date"].isoformat()])
    metadata.append(["Grouping", request.GET.get("group_by") or "major"])
    metadata.append(["Generated At", timezone.now().isoformat()])
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="Mining360_Parts_Classification.xlsx"'
    return response


@login_required
@require_http_methods(["POST"])
def command_center_mark_reviewed_api(request):
    if not _command_center_allowed(request.user):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    filter_hash = str(_payload(request).get("filter_hash") or "").strip()
    visit = BusinessCommandCenterUserVisit.objects.filter(user=request.user, filter_hash=filter_hash).first()
    if not visit:
        return JsonResponse({"detail": "Visit context not found."}, status=404)
    visit.marked_reviewed_at = timezone.now()
    visit.save(update_fields=["marked_reviewed_at"])
    return JsonResponse({"ok": True, "marked_reviewed_at": visit.marked_reviewed_at})


@login_required
@require_http_methods(["GET", "POST", "DELETE"])
def command_center_watchlist_api(request):
    if not _command_center_allowed(request.user) or not feature_enabled("ENABLE_BUSINESS_COMMAND_CENTER_WATCHLIST", request.user):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    if request.method == "POST":
        data = _payload(request)
        entity_type_input = str(data.get("entity_type") or "").strip().lower().replace("_", " ")
        entity_type = {"customer": "Customer", "country": "Country", "key account": "Key Account", "business line": "Business Line"}.get(entity_type_input)
        entity_id = str(data.get("entity_id") or "").strip()
        if not entity_type or not entity_id:
            return JsonResponse({"detail": "A valid watchlist entity is required."}, status=400)
        item, _ = BusinessCommandCenterWatchlist.objects.update_or_create(
            user=request.user, entity_type=entity_type, entity_id=entity_id,
            defaults={"display_name": str(data.get("display_name") or entity_id)[:255], "active": True},
        )
        return JsonResponse({"ok": True, "id": str(item.id)}, status=201)
    if request.method == "DELETE":
        item_id = str(_payload(request).get("id") or "").strip()
        updated = BusinessCommandCenterWatchlist.objects.filter(user=request.user, id=item_id).update(active=False)
        return JsonResponse({"ok": bool(updated)}, status=200 if updated else 404)
    rows = BusinessCommandCenterWatchlist.objects.filter(user=request.user, active=True)
    return JsonResponse({"results": [{
        "id": str(item.id), "entity_type": item.entity_type, "entity_id": item.entity_id,
        "display_name": item.display_name,
    } for item in rows]})


@login_required
def command_center_export_api(request):
    if not _command_center_allowed(request.user) or not has_business_review_permission(request.user, "export_business_command_center"):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    try:
        payload = BusinessCommandCenterService(request.user, request.GET).bootstrap()
    except BusinessCommandCenterInputError as exc:
        return JsonResponse({"detail": str(exc)}, status=400)
    metrics = {
        "mining_revenue_ytd_eur": payload["hero"]["revenue"],
        "mining_revenue_previous_year_eur": payload["hero"]["comparison_revenue"],
        "unallocated_revenue_eur": payload["confidence"]["unallocated_revenue"],
        "revenue_coverage_pct": payload["confidence"]["customer_coverage"],
        "revenue_by_lob": {item["label"]: item["revenue"] for item in payload["business_lines"]},
    }
    content = BusinessReviewExportService.build(
        context={**payload["context"], **payload["freshness"]},
        confidence=payload["confidence"], metrics=metrics, portfolio=[],
    )
    response = HttpResponse(content, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="Mining360_Business_Command_Center.xlsx"'
    return response


def _payload(request):
    try:
        return json.loads(request.body or b"{}")
    except (TypeError, json.JSONDecodeError):
        return {}


def _audit_value(value):
    return json.loads(json.dumps(value, cls=DjangoJSONEncoder))


def _lens(request):
    value = request.GET.get("lens", "ALL").strip().upper()
    return value if value in {"ALL", *MINING_REVENUE_LOBS} else "ALL"


def _review_context(request):
    publication = MappingPublication.objects.filter(status="Published").order_by("-version").first()
    if not publication:
        return None, [], None
    snapshot = BusinessReviewSnapshotService.generate(request.user, publication)
    all_rows = list((publication.snapshot_json or {}).get("mappings", []))
    rows = filter_published_rows(all_rows, request.user)
    account_id = request.GET.get("account", "").strip()
    minesite_id = request.GET.get("minesite", "").strip()
    country = request.GET.get("country", "").strip().casefold()
    if account_id:
        rows = [row for row in rows if row.get("account_id") == account_id]
    if minesite_id:
        rows = [row for row in rows if row.get("minesite_id") == minesite_id]
    if country:
        site_ids = set(MineSite.objects.filter(country__iexact=country).values_list("id", flat=True))
        rows = [
            row for row in rows
            if row.get("minesite_id") in {str(value) for value in site_ids}
            or any(str(record.get("country") or "").strip().casefold() == country for record in row.get("source_records", []))
        ]
    return snapshot, rows, publication


def _scope_is_complete(snapshot, rows):
    return len({row.get("mapping_id") for row in rows}) == snapshot.mapping_publication.mapping_count


def _confidence(snapshot, publication, rows, metrics):
    values = dict(metrics)
    if not _scope_is_complete(snapshot, rows):
        for key in ("account_coverage_pct", "revenue_coverage_pct", "fleet_coverage_pct", "unallocated_revenue_eur"):
            values[key] = None
    return BusinessReviewDataConfidenceService.evaluate(
        publication=publication,
        source_run=snapshot.source_synchronization,
        metrics=values,
    )


def _filter_options(publication, user):
    rows = filter_published_rows(list((publication.snapshot_json or {}).get("mappings", [])), user)
    account_options = {}
    minesite_options = {}
    countries = set()
    site_ids = {row.get("minesite_id") for row in rows if row.get("minesite_id")}
    site_countries = {
        str(item["id"]): item["country"]
        for item in MineSite.objects.filter(id__in=site_ids).values("id", "country")
    }
    for row in rows:
        if row.get("account_id"):
            account_options[row["account_id"]] = row.get("account_name") or row.get("account_code")
        if row.get("minesite_id"):
            minesite_options[row["minesite_id"]] = row.get("minesite_name")
            if site_countries.get(row["minesite_id"]):
                countries.add(site_countries[row["minesite_id"]])
        countries.update(record.get("country") for record in row.get("source_records", []) if record.get("country"))
    return {
        "countries": sorted(countries),
        "accounts": [{"id": key, "name": value} for key, value in sorted(account_options.items(), key=lambda item: item[1])],
        "minesites": [{"id": key, "name": value} for key, value in sorted(minesite_options.items(), key=lambda item: item[1])],
    }


def _scoped_changes(publication, rows):
    allowed = {row.get("mapping_id") for row in rows}
    changes = publication.change_summary_json or {}
    return {
        key: [item for item in changes.get(key, []) if item.get("mapping_id") in allowed]
        for key in ("added", "changed", "removed")
    }


def _scoped_metrics(snapshot, rows, lens):
    source_codes = {code for row in rows for code in row.get("source_account_codes", []) if code}
    site_names = {normalize_business_name(row.get("minesite_name")) for row in rows if row.get("minesite_name")}
    revenue = RevenueSourceSnapshot.objects.filter(
        active=True, division__iexact=MINING_DIVISION, lob__in=MINING_REVENUE_LOBS,
        source_account_code__in=source_codes,
    )
    lens_revenue = revenue.filter(lob=lens) if lens in MINING_REVENUE_LOBS else revenue
    ytd = lens_revenue.aggregate(value=Sum("revenue_ytd_eur"))["value"] or Decimal("0")
    previous = lens_revenue.aggregate(value=Sum("revenue_previous_year_eur"))["value"] or Decimal("0")
    by_lob = {item["lob"]: float(item["value"] or 0) for item in revenue.values("lob").annotate(value=Sum("revenue_ytd_eur"))}
    fleet = EquipmentFleetAnalysis.objects.filter(active=True, normalized_site__in=site_names)
    fleet_count = fleet.count()
    return {
        "mining_revenue_ytd_eur": float(ytd),
        "mining_revenue_previous_year_eur": float(previous),
        "revenue_change_pct": round(float((ytd - previous) / previous * 100), 1) if previous else None,
        "revenue_by_lob": {lob: by_lob.get(lob, 0) for lob in MINING_REVENUE_LOBS},
        "fleet_count": fleet_count,
        "revenue_per_equipment_eur": float(ytd / fleet_count) if fleet_count else None,
        "published_account_count": len({row.get("account_id") for row in rows if row.get("account_id")}),
        "published_minesite_count": len(site_names),
        "revenue_assigned_eur": float(ytd),
        "unallocated_revenue_eur": snapshot.metrics_json.get("unallocated_revenue_eur") if _scope_is_complete(snapshot, rows) else None,
        "revenue_coverage_pct": snapshot.metrics_json.get("revenue_coverage_pct") if _scope_is_complete(snapshot, rows) else None,
        "fleet_coverage_pct": snapshot.metrics_json.get("fleet_coverage_pct") if _scope_is_complete(snapshot, rows) else None,
        "account_coverage_pct": snapshot.metrics_json.get("account_coverage_pct") if _scope_is_complete(snapshot, rows) else None,
    }


def _not_ready():
    return JsonResponse({
        "ready": False,
        "status": "NOT_READY",
        "message": "Business Overview is not ready yet. A Published Mapping version is required before managerial analysis can be generated.",
    })


@login_required
def business_review(request):
    if not business_review_enabled(request.user):
        return HttpResponseForbidden("You do not have access to Business Overview.")
    return render(request, "reports/business_review.html", {
        "active_section": "business-review",
        "can_create_action": has_business_review_permission(request.user, "create_business_review_action") and feature_enabled("ENABLE_BUSINESS_REVIEW_ACTIONS", request.user),
        "can_export": has_business_review_permission(request.user, "export_business_review") and feature_enabled("ENABLE_BUSINESS_REVIEW_EXPORT", request.user),
        "can_preview": has_business_review_permission(request.user, "preview_business_review_draft"),
        "feature_flags": {
            "portfolio": feature_enabled("ENABLE_BUSINESS_REVIEW_PORTFOLIO", request.user),
            "opportunities": feature_enabled("ENABLE_BUSINESS_OPPORTUNITY_ENGINE", request.user),
            "risks": feature_enabled("ENABLE_BUSINESS_RISK_ENGINE", request.user),
            "actions": feature_enabled("ENABLE_BUSINESS_REVIEW_ACTIONS", request.user),
            "decisions": feature_enabled("ENABLE_BUSINESS_REVIEW_DECISIONS", request.user),
            "export": feature_enabled("ENABLE_BUSINESS_REVIEW_EXPORT", request.user),
        },
    })


@login_required
def context_api(request):
    if not business_review_enabled(request.user):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    snapshot, rows, publication = _review_context(request)
    if not snapshot:
        return _not_ready()
    context = dict(snapshot.context_json)
    context.update({"scope": "authorized", "revenue_lens": _lens(request), "published_mapping_version": publication.version})
    return JsonResponse({"ready": True, "snapshot_id": str(snapshot.id), "context": context, "warnings": snapshot.warnings_json})


@login_required
def overview_api(request):
    if not business_review_enabled(request.user):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    snapshot, rows, publication = _review_context(request)
    if not snapshot:
        return _not_ready()
    lens = _lens(request)
    metrics = _scoped_metrics(snapshot, rows, lens)
    confidence = _confidence(snapshot, publication, rows, metrics)
    BusinessRiskService.regenerate(snapshot)
    allowed_accounts = {row.get("account_id") for row in rows}
    allowed_sites = {row.get("minesite_id") for row in rows}
    risk_scope = Q(business_account_id__in=allowed_accounts) | Q(minesite_id__in=allowed_sites)
    if _scope_is_complete(snapshot, rows):
        risk_scope |= Q(business_account__isnull=True, minesite__isnull=True)
    risks = list(snapshot.risks.filter(status="Open").filter(risk_scope).values(
        "id", "risk_code", "severity", "business_impact_json"
    )[:5])
    actions = list(snapshot.management_actions.exclude(status__in=["Completed", "Cancelled"]).values("id", "title", "priority", "status", "due_date")[:5])
    return JsonResponse({
        "ready": True,
        "context": {**snapshot.context_json, "scope": "authorized", "revenue_lens": lens, "published_mapping_version": publication.version},
        "freshness": {"revenue_snapshot_at": snapshot.context_json.get("revenue_snapshot_at"), "fleet_snapshot_at": snapshot.context_json.get("fleet_snapshot_at"), "review_generated_at": snapshot.generated_at},
        "confidence": confidence,
        "metrics": metrics,
        "changes": _scoped_changes(publication, rows),
        "filter_options": _filter_options(publication, request.user),
        "attention_items": risks,
        "risks": risks,
        "actions": actions,
    })


@login_required
def portfolio_api(request):
    if not business_review_enabled(request.user) or not feature_enabled("ENABLE_BUSINESS_REVIEW_PORTFOLIO", request.user):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    snapshot, rows, _ = _review_context(request)
    if not snapshot:
        return _not_ready()
    lens = _lens(request)
    allowed_sites = {row.get("minesite_id") for row in rows}
    portfolio = [item for item in BusinessReviewSnapshotService.site_portfolio(snapshot, lens) if item.get("entity_id") in allowed_sites]
    portfolio, thresholds = BusinessOpportunityRuleEngine.regenerate(snapshot, portfolio, lens)
    return JsonResponse({"ready": True, "snapshot_id": str(snapshot.id), "lens": lens, "thresholds": thresholds, "results": portfolio})


@login_required
def opportunities_api(request):
    if not business_review_enabled(request.user) or not feature_enabled("ENABLE_BUSINESS_OPPORTUNITY_ENGINE", request.user):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    snapshot, rows, _ = _review_context(request)
    if not snapshot:
        return _not_ready()
    allowed_sites = {row.get("minesite_id") for row in rows}
    queryset = snapshot.opportunities.filter(minesite_id__in=allowed_sites, revenue_lens=_lens(request)).select_related("minesite")
    return JsonResponse({"ready": True, "results": [{
        "id": str(item.id), "code": item.opportunity_code, "classification": item.classification,
        "severity": item.severity, "status": item.status, "minesite": item.minesite.canonical_minesite_name if item.minesite else None,
        "metrics": item.metric_values_json, "thresholds": item.threshold_values_json, "evidence": item.evidence_json,
    } for item in queryset[:200]]})


@login_required
def risks_api(request):
    if not business_review_enabled(request.user) or not feature_enabled("ENABLE_BUSINESS_RISK_ENGINE", request.user):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    snapshot, rows, _ = _review_context(request)
    if not snapshot:
        return _not_ready()
    BusinessRiskService.regenerate(snapshot)
    allowed_accounts = {row.get("account_id") for row in rows}
    allowed_sites = {row.get("minesite_id") for row in rows}
    scope_filter = Q(business_account_id__in=allowed_accounts) | Q(minesite_id__in=allowed_sites)
    if _scope_is_complete(snapshot, rows):
        scope_filter |= Q(business_account__isnull=True, minesite__isnull=True)
    return JsonResponse({"ready": True, "results": [{
        "id": str(item.id), "code": item.risk_code, "severity": item.severity, "status": item.status,
        "impact": item.business_impact_json, "evidence": item.evidence_json, "detected_at": item.detected_at,
    } for item in snapshot.risks.filter(scope_filter)[:200]]})


@login_required
def accounts_api(request):
    if not business_review_enabled(request.user):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    snapshot, rows, _ = _review_context(request)
    if not snapshot:
        return _not_ready()
    grouped = {}
    for row in rows:
        item = grouped.setdefault(row["account_id"], {"id": row["account_id"], "name": row["account_name"], "code": row["account_code"], "minesites": [], "roles": set(), "source_record_count": len(row.get("source_records", []))})
        if row.get("minesite_name") and row["minesite_name"] not in item["minesites"]:
            item["minesites"].append(row["minesite_name"])
        item["roles"].add(row.get("role"))
    results = [{**item, "roles": sorted(role for role in item["roles"] if role)} for item in grouped.values()]
    search = request.GET.get("search", "").strip().casefold()
    if search:
        results = [item for item in results if search in " ".join([
            str(item.get("name") or ""), str(item.get("code") or ""), *item.get("minesites", [])
        ]).casefold()]
    return JsonResponse({"ready": True, "count": len(results), "results": sorted(results, key=lambda item: item["name"])})


@login_required
def account_detail_api(request, account_id):
    if not business_review_enabled(request.user):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    snapshot, rows, publication = _review_context(request)
    scoped = [row for row in rows if row.get("account_id") == str(account_id)]
    if not snapshot or not scoped:
        return JsonResponse({"detail": "Not found"}, status=404)
    source_codes = {code for row in scoped for code in row.get("source_account_codes", [])}
    revenue = RevenueSourceSnapshot.objects.filter(active=True, division__iexact=MINING_DIVISION, source_account_code__in=source_codes, lob__in=MINING_REVENUE_LOBS)
    by_lob = {item["lob"]: float(item["value"] or 0) for item in revenue.values("lob").annotate(value=Sum("revenue_ytd_eur"))}
    sites = [row.get("minesite_name") for row in scoped if row.get("minesite_name")]
    normalized_sites = [normalize_business_name(site) for site in sites]
    fleet = EquipmentFleetAnalysis.objects.filter(active=True, normalized_site__in=normalized_sites)
    return JsonResponse({"ready": True, "account": {"id": str(account_id), "name": scoped[0]["account_name"], "code": scoped[0]["account_code"]}, "relationships": scoped, "revenue": {"total": sum(by_lob.values()), "by_lob": by_lob}, "fleet": {"count": fleet.count(), "serial_count": fleet.exclude(serial_number="").values("serial_number").distinct().count(), "models": list(fleet.exclude(model="").values_list("model", flat=True).distinct()[:50])}, "published_mapping_version": publication.version, "data_confidence": snapshot.confidence_status})


@login_required
def minesites_api(request):
    if not business_review_enabled(request.user):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    snapshot, rows, _ = _review_context(request)
    if not snapshot:
        return _not_ready()
    portfolio = BusinessReviewSnapshotService.site_portfolio(snapshot, _lens(request))
    allowed = {row.get("minesite_id") for row in rows}
    results = [item for item in portfolio if item.get("entity_id") in allowed]
    search = request.GET.get("search", "").strip().casefold()
    if search:
        results = [item for item in results if search in str(item.get("name") or "").casefold()]
    return JsonResponse({"ready": True, "results": results})


@login_required
def minesite_detail_api(request, minesite_id):
    if not business_review_enabled(request.user):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    snapshot, rows, publication = _review_context(request)
    scoped = [row for row in rows if row.get("minesite_id") == str(minesite_id)]
    if not snapshot or not scoped:
        return JsonResponse({"detail": "Not found"}, status=404)
    portfolio = next((item for item in BusinessReviewSnapshotService.site_portfolio(snapshot, _lens(request)) if item.get("entity_id") == str(minesite_id)), None)
    site = MineSite.objects.filter(pk=minesite_id).first()
    fleet = EquipmentFleetAnalysis.objects.filter(active=True, normalized_site=normalize_business_name(scoped[0]["minesite_name"]))
    return JsonResponse({"ready": True, "minesite": {"id": str(minesite_id), "name": scoped[0]["minesite_name"], "country": site.country if site else None, "owner": site.owner if site else None, "operator": site.operator if site else None}, "account_network": scoped, "portfolio": portfolio, "fleet": {"count": fleet.count(), "serial_count": fleet.exclude(serial_number="").values("serial_number").distinct().count(), "models": list(fleet.exclude(model="").values_list("model", flat=True).distinct()[:50]), "families": list(fleet.exclude(equipment_family="").values_list("equipment_family", flat=True).distinct()[:50])}, "published_mapping_version": publication.version, "data_confidence": snapshot.confidence_status})


@login_required
def data_confidence_api(request):
    if not has_business_review_permission(request.user, "view_business_review_data_confidence"):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    snapshot, rows, publication = _review_context(request)
    if not snapshot:
        return _not_ready()
    metrics = _scoped_metrics(snapshot, rows, _lens(request))
    return JsonResponse(_confidence(snapshot, publication, rows, metrics))


@login_required
def changes_api(request):
    if not business_review_enabled(request.user):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    _, rows, publication = _review_context(request)
    return _not_ready() if not publication else JsonResponse({"ready": True, "published_mapping_version": publication.version, "changes": _scoped_changes(publication, rows)})


@login_required
def comparison_api(request):
    if not has_business_review_permission(request.user, "compare_business_entities"):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    snapshot, rows, _ = _review_context(request)
    if not snapshot:
        return _not_ready()
    requested = {value for value in request.GET.getlist("minesite") if value}
    portfolio = BusinessReviewSnapshotService.site_portfolio(snapshot, _lens(request))
    allowed = {row.get("minesite_id") for row in rows}
    return JsonResponse({"ready": True, "results": [item for item in portfolio if item.get("entity_id") in allowed and (not requested or item.get("entity_id") in requested)]})


@login_required
@require_http_methods(["GET", "POST"])
def actions_api(request):
    if not business_review_enabled(request.user) or not feature_enabled("ENABLE_BUSINESS_REVIEW_ACTIONS", request.user):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    snapshot, rows, _ = _review_context(request)
    if not snapshot:
        return _not_ready()
    allowed_accounts = {row.get("account_id") for row in rows}
    allowed_sites = {row.get("minesite_id") for row in rows}
    if request.method == "POST":
        if not has_business_review_permission(request.user, "create_business_review_action"):
            return JsonResponse({"detail": "Forbidden"}, status=403)
        data = _payload(request)
        title = str(data.get("title") or "").strip()
        if not title:
            return JsonResponse({"detail": "Action title is required."}, status=400)
        account_id = data.get("business_account_id") or None
        site_id = data.get("minesite_id") or None
        if (account_id and account_id not in allowed_accounts) or (site_id and site_id not in allowed_sites):
            return JsonResponse({"detail": "Forbidden"}, status=403)
        if data.get("owner_id") and not has_business_review_permission(request.user, "assign_business_review_action"):
            return JsonResponse({"detail": "You do not have permission to assign this action."}, status=403)
        if data.get("owner_id") and not get_user_model().objects.filter(pk=data["owner_id"], is_active=True).exists():
            return JsonResponse({"detail": "The selected owner is not available."}, status=400)
        action = BusinessReviewAction.objects.create(
            snapshot=snapshot, title=title, description=str(data.get("description") or ""),
            action_type=str(data.get("action_type") or "Management Review"), business_account_id=account_id,
            minesite_id=site_id, priority=data.get("priority") if data.get("priority") in dict(BusinessReviewAction.PRIORITIES) else "Medium",
            due_date=data.get("due_date") or None, status="Assigned" if data.get("owner_id") else "Open",
            owner_id=data.get("owner_id") or None, business_lens=data.get("business_lens") or _lens(request), created_by=request.user,
        )
        MappingAuditLog.objects.create(
            actor=request.user, action="BUSINESS_REVIEW_ACTION_CREATED", entity_type="BusinessReviewAction",
            entity_id=str(action.id), new_value_json=_audit_value(_action_json(action)),
            source_ip=request.META.get("REMOTE_ADDR"),
        )
        return JsonResponse({"ok": True, "action": _action_json(action)}, status=201)
    queryset = snapshot.management_actions.filter(Q(business_account_id__in=allowed_accounts) | Q(minesite_id__in=allowed_sites) | Q(business_account__isnull=True, minesite__isnull=True)).select_related("owner", "business_account", "minesite")
    return JsonResponse({"ready": True, "results": [_action_json(item) for item in queryset[:500]]})


def _action_json(item):
    return {"id": str(item.id), "title": item.title, "description": item.description, "priority": item.priority, "status": item.status, "owner": item.owner.get_full_name() or item.owner.get_username() if item.owner else None, "due_date": item.due_date, "account": item.business_account.canonical_account_name if item.business_account else None, "minesite": item.minesite.canonical_minesite_name if item.minesite else None, "updated_at": item.updated_at}


@login_required
@require_http_methods(["PATCH"])
def action_detail_api(request, action_id):
    if not business_review_enabled(request.user) or not feature_enabled("ENABLE_BUSINESS_REVIEW_ACTIONS", request.user):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    snapshot, rows, _ = _review_context(request)
    allowed_accounts = {row.get("account_id") for row in rows}
    allowed_sites = {row.get("minesite_id") for row in rows}
    action = BusinessReviewAction.objects.filter(pk=action_id, snapshot=snapshot).first() if snapshot else None
    if not action:
        return JsonResponse({"detail": "Not found"}, status=404)
    if (action.business_account_id and str(action.business_account_id) not in allowed_accounts) or (action.minesite_id and str(action.minesite_id) not in allowed_sites):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    can_manage = action.created_by_id == request.user.id or action.owner_id == request.user.id or has_business_review_permission(request.user, "assign_business_review_action")
    if not can_manage:
        return JsonResponse({"detail": "Forbidden"}, status=403)
    data = _payload(request)
    previous = _action_json(action)
    if data.get("status") == "Completed" and not has_business_review_permission(request.user, "complete_business_review_action"):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    for field in ("title", "description", "priority", "status", "due_date", "notes"):
        if field in data:
            setattr(action, field, data[field] or (None if field == "due_date" else ""))
    if action.status == "Completed" and not action.completed_at:
        action.completed_at = timezone.now()
    action.save()
    MappingAuditLog.objects.create(
        actor=request.user, action="BUSINESS_REVIEW_ACTION_UPDATED", entity_type="BusinessReviewAction",
        entity_id=str(action.id), previous_value_json=_audit_value(previous), new_value_json=_audit_value(_action_json(action)),
        source_ip=request.META.get("REMOTE_ADDR"),
    )
    return JsonResponse({"ok": True, "action": _action_json(action)})


@login_required
@require_http_methods(["GET", "POST"])
def decisions_api(request):
    if not business_review_enabled(request.user) or not feature_enabled("ENABLE_BUSINESS_REVIEW_DECISIONS", request.user):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    snapshot, rows, _ = _review_context(request)
    if not snapshot:
        return _not_ready()
    if request.method == "POST":
        if not has_business_review_permission(request.user, "create_business_decision"):
            return JsonResponse({"detail": "Forbidden"}, status=403)
        data = _payload(request)
        account_id = data.get("business_account_id") or None
        site_id = data.get("minesite_id") or None
        allowed_accounts = {row.get("account_id") for row in rows}
        allowed_sites = {row.get("minesite_id") for row in rows}
        if (account_id and account_id not in allowed_accounts) or (site_id and site_id not in allowed_sites):
            return JsonResponse({"detail": "Forbidden"}, status=403)
        decision = BusinessDecision.objects.create(snapshot=snapshot, title=data.get("title", ""), statement=data.get("statement", ""), action_id=data.get("action_id") or None, business_account_id=account_id, minesite_id=site_id, decision_date=data.get("decision_date") or timezone.localdate(), effective_date=data.get("effective_date") or None, review_date=data.get("review_date") or None, expected_result=data.get("expected_result", ""), evidence_json=data.get("evidence", []), decided_by=request.user)
        return JsonResponse({"ok": True, "id": str(decision.id)}, status=201)
    allowed_accounts = {row.get("account_id") for row in rows}
    allowed_sites = {row.get("minesite_id") for row in rows}
    decision_scope = Q(business_account_id__in=allowed_accounts) | Q(minesite_id__in=allowed_sites)
    if _scope_is_complete(snapshot, rows):
        decision_scope |= Q(business_account__isnull=True, minesite__isnull=True)
    return JsonResponse({"ready": True, "results": list(snapshot.decisions.filter(decision_scope).values("id", "title", "statement", "decision_date", "review_date", "status")[:500])})


@login_required
@require_http_methods(["GET", "POST"])
def saved_views_api(request):
    if not business_review_enabled(request.user):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    if request.method == "POST":
        data = _payload(request)
        view, _ = BusinessReviewSavedView.objects.update_or_create(user=request.user, name=data.get("name", "My Business Overview"), defaults={"filters_json": data.get("filters", {}), "sorting_json": data.get("sorting", []), "visualization": data.get("visualization", "executive_overview")})
        return JsonResponse({"ok": True, "id": str(view.id)}, status=201)
    return JsonResponse({"results": list(request.user.business_review_saved_views.values("id", "name", "filters_json", "sorting_json", "visualization", "updated_at"))})


@login_required
def export_api(request):
    if not has_business_review_permission(request.user, "export_business_review") or not feature_enabled("ENABLE_BUSINESS_REVIEW_EXPORT", request.user):
        return JsonResponse({"detail": "Forbidden"}, status=403)
    snapshot, rows, publication = _review_context(request)
    if not snapshot:
        return _not_ready()
    lens = _lens(request)
    metrics = _scoped_metrics(snapshot, rows, lens)
    confidence = _confidence(snapshot, publication, rows, metrics)
    allowed_sites = {row.get("minesite_id") for row in rows}
    portfolio = [item for item in BusinessReviewSnapshotService.site_portfolio(snapshot, lens) if item.get("entity_id") in allowed_sites]
    portfolio, _ = BusinessOpportunityRuleEngine.classify(portfolio, lens)
    content = BusinessReviewExportService.build(context={**snapshot.context_json, "revenue_lens": lens, "published_mapping_version": publication.version}, confidence=confidence, metrics=metrics, portfolio=portfolio)
    response = HttpResponse(content, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Mining360_Business_Review_v{publication.version}_{lens}.xlsx"'
    return response
