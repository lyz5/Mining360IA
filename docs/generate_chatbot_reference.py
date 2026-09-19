"""Generate the governed Mining360 chatbot reference document."""

from __future__ import annotations

import os
import sys
from datetime import date
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "Mining360IA.settings")

import django

django.setup()

from django.conf import settings
from django.db.models import Count
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from reports.fleet_performance_intelligence_service import CORE_METRICS, METRIC_BUNDLES
from reports.machine_performance_intent_service import INTENT_TYPES
from reports.models import (
    AIAnswerabilityConfiguration,
    AIAgent,
    AIAgentCapability,
    AIDaxTemplate,
    AIFilterMapping,
    AIIntentResponseTemplateMapping,
    AIMetricMapping,
    AIQuestionExample,
    AIResponseTemplate,
    BusinessDataField,
    BusinessPerformanceConfig,
    BusinessPerformanceMapping,
    KnowledgeSynonym,
    PowerBIReport,
    AICapabilityOperation,
    AIChatSuggestion,
    AIActionContract,
    AIDependencyHealthSnapshot,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Mining360_Chatbot_Reference_Complete_2026-09-03.docx"
NAVY = "17213B"
TEAL = "2D5753"
YELLOW = "F5C400"
LIGHT = "F3F5F8"
MID = "D9E0E8"
RED = "A53020"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")



def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, end))


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.7)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color in (("Title", 30, NAVY), ("Heading 1", 20, NAVY), ("Heading 2", 14, TEAL), ("Heading 3", 11, NAVY)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if name != "Heading 1" else 18)
        style.paragraph_format.space_after = Pt(5)
    styles["Heading 1"].paragraph_format.page_break_before = True

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = "Consolas"
        code.font.size = Pt(8)
        code.font.color.rgb = RGBColor.from_string(NAVY)
        code.paragraph_format.left_indent = Cm(0.4)
        code.paragraph_format.right_indent = Cm(0.4)
        code.paragraph_format.space_before = Pt(4)
        code.paragraph_format.space_after = Pt(6)

    header = section.header.paragraphs[0]
    header.text = "MINING 360  |  CHATBOT - REFERENCE FONCTIONNELLE ET TECHNIQUE"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(7.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(TEAL)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Confidentiel interne  |  Mining360  |  ")
    add_field(footer, "PAGE")


def add_title(doc):
    for _ in range(3):
        doc.add_paragraph()
    eyebrow = doc.add_paragraph()
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = eyebrow.add_run("NEEMBA  |  MINING 360")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(TEAL)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Chatbot Mining 360")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Reference fonctionnelle, technique, securite et exploitation")
    run.font.size = Pt(17)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = line.add_run("━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━")
    run.font.color.rgb = RGBColor.from_string(YELLOW)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Version 1.1\nEtat de reference : 3 septembre 2026\nEnvironnement analyse : Mining360IA / configuration locale active").bold = True
    doc.add_paragraph()
    note = doc.add_table(rows=1, cols=1)
    note.alignment = WD_TABLE_ALIGNMENT.CENTER
    note.autofit = False
    note.columns[0].width = Cm(15)
    cell = note.cell(0, 0)
    set_cell_shading(cell, LIGHT)
    set_cell_margins(cell, 180, 220, 180, 220)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(
        "Ce document decrit l'implementation reelle observee dans le code et la base de configuration. "
        "Il distingue les fonctions disponibles, limitees, sous feature flag et non configurees."
    )
    r.italic = True
    doc.add_page_break()


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_p(doc, text, *, bold_prefix=""):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(item, tuple):
            p.add_run(item[0]).bold = True
            p.add_run(item[1])
        else:
            p.add_run(str(item))


def add_numbered(doc, items):
    for item in items:
        doc.add_paragraph(str(item), style="List Number")


def add_callout(doc, title, text, color=TEAL):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT)
    set_cell_margins(cell, 140, 180, 140, 180)
    p = cell.paragraphs[0]
    r = p.add_run(title + "\n")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(color)
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code(doc, text):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "EEF1F5")
    set_cell_margins(cell, 120, 140, 120, 140)
    p = cell.paragraphs[0]
    p.style = doc.styles["Code Block"]
    p.add_run(text)


def add_table(doc, headers, rows, widths=None, font_size=7.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = str(header)
        set_cell_shading(cell, TEAL)
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(font_size)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = "" if value is None else str(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[index])
            if row_index % 2:
                set_cell_shading(cells[index], "F8F9FB")
            for paragraph in cells[index].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = RGBColor.from_string(NAVY)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def bool_text(value):
    return "Oui" if value else "Non"


def list_text(value):
    if not value:
        return "-"
    return ", ".join(str(item) for item in value)


def feature_flags():
    names = [
        "ENABLE_PERSISTENT_CONVERSATIONS", "ENABLE_CONVERSATION_ARCHIVE",
        "ENABLE_CONVERSATION_RENAME", "ENABLE_CONVERSATION_ARTIFACTS",
        "ENABLE_ADAPTIVE_PERFORMANCE_RESPONSES", "ENABLE_FLEET_INVENTORY_CHAT",
        "ENABLE_FLEET_EXCEL_EXPORT", "ENABLE_EQUIPMENT_SERIAL_LOOKUP",
        "ENABLE_EQUIPMENT_CODE_LOOKUP", "ENABLE_FLEET_FUZZY_SITE_MATCHING",
        "ENABLE_COMPLETE_FLEET_PERFORMANCE_CHAT", "ENABLE_FLEET_PERFORMANCE_COMPARISON",
        "ENABLE_FLEET_PERFORMANCE_TRENDS", "ENABLE_FLEET_PERFORMANCE_RANKING",
        "ENABLE_PLANNED_UNPLANNED_ANALYSIS", "ENABLE_FLEET_PERFORMANCE_EXPORT",
        "ENABLE_FLEET_PERFORMANCE_DIAGNOSTICS", "ENABLE_CONVERSATIONAL_FOLLOW_UP_RESOLUTION",
        "ENABLE_CHATBOT_CAPABILITY_DISCOVERY", "ENABLE_ANSWERABILITY_GUARD",
        "ENABLE_GROUNDED_RESPONSE_GUARD", "ENABLE_GRACEFUL_ABSTENTION",
        "ENABLE_DATA_GAP_REGISTRY", "ENABLE_CONTEXTUAL_CAPABILITY_HELP",
        "ENABLE_USER_OWNS_DATA_EMBEDDING", "ENABLE_DELEGATED_TOKEN_REFRESH",
        "ENABLE_DOWNTIME_MAPPING_CHECK", "ENABLE_FULL_AI_DOWNTIME_AUDIT",
        "ENABLE_DOWNTIME_MAPPING_WRITEBACK", "ENABLE_DOWNTIME_MAPPING_BATCH_PROCESSING",
        "ENABLE_DOWNTIME_MAPPING_KNOWLEDGE_LEARNING",
        "ENABLE_CERTIFIED_CHAT_SUGGESTIONS", "ENABLE_OPERATION_LEVEL_READINESS",
        "ENABLE_ACTION_ELIGIBILITY", "ENABLE_CHAT_PRODUCTION_STATE_MACHINE",
        "ENABLE_CHAT_DEPENDENCY_HEALTH", "ENABLE_CHAT_READINESS_DASHBOARD",
        "ENABLE_CHAT_DEAD_SUGGESTION_AUTO_HIDE", "ENABLE_REAL_PILOT_COHORTS",
    ]
    return [(name, getattr(settings, name, "Non defini")) for name in names]


def document_control(doc):
    add_heading(doc, "Controle du document")
    add_table(doc, ["Champ", "Valeur"], [
        ("Titre", "Chatbot Mining 360 - Reference complete"),
        ("Version", "1.1"),
        ("Date de reference", "3 septembre 2026"),
        ("Source", "Code Mining360IA, migrations et base de configuration active"),
        ("Audience", "Metiers, Data/BI, IT, support, securite, administrateurs Mining360"),
        ("Classification", "Confidentiel interne"),
        ("Tests observes", "487 tests reports et 30 tests de deploiement valides; 86 tests chatbot cibles valides"),
    ], widths=[4, 12], font_size=8)
    add_heading(doc, "Sommaire", 2)
    p = doc.add_paragraph()
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u')
    add_p(doc, "Dans Word, utilisez Mettre a jour la table pour recalculer le sommaire et les numeros de page.")


def executive_summary(doc):
    add_heading(doc, "1. Synthese executive")
    add_p(doc, "Mining 360 AI est un chatbot analytique gouverne. Il combine des traitements deterministes, des requetes semantiques Power BI, une base de connaissances validee et, lorsque necessaire, un fournisseur d'IA pour la formulation ou l'analyse textuelle. Les chiffres, machines, sites et faits internes ne doivent pas etre inventes par le modele de langage.")
    add_heading(doc, "Capacites principales", 2)
    add_bullets(doc, [
        ("Flotte et equipements. ", "Inventaire complet d'un MineSite, regroupement par modele, recherche par numero de serie ou code equipement, details machine et export Excel."),
        ("Performance et fiabilite. ", "Physical Availability, MTBS, MTBF, MTTR, downtime planifie et non planifie, avec valeur, vue multi-KPI, tendance, comparaison et classement."),
        ("Downtime et causes. ", "Drivers, Pareto, equipements affectes, evenements et composants. Les analyses IA avancees ne sont annoncees que lorsqu'elles sont certifiees."),
        ("Parts Sales. ", "Chiffre d'affaires Parts YTD, global, par client ou territoire/MineSite, sur les mesures de facture officielles."),
        ("Knowledge et Best Practices. ", "Definitions KPI, recherche documentaire validee, recommandations et citations de sources."),
        ("Reporting. ", "Recherche et ouverture de rapports, navigation Power BI et application du contexte valide."),
        ("Conversation. ", "Francais/anglais, follow-ups, contexte persistant, aide personnalisee, abstention explicite et retry controle."),
    ])
    add_callout(doc, "Principe de confiance", "Les donnees structurees viennent des sources configurees. Le backend construit et valide le plan. Le fournisseur d'IA peut expliquer le resultat mais ne cree pas les valeurs operationnelles.")
    add_heading(doc, "Etat de disponibilite actuel", 2)
    caps = AIAgentCapability.objects.select_related("agent").order_by("display_order", "display_name")
    add_table(doc, ["Capacite", "Agent", "Etat", "Operations"], [
        (c.display_name, c.agent.name, c.readiness_status, list_text(c.supported_operations_json))
        for c in caps if c.readiness_status in {"Ready", "Limited"}
    ], widths=[4.2, 3.5, 2.2, 7], font_size=7.2)
    add_p(doc, "Les capacites historiques marquees Needs Configuration restent dans le registre administratif mais ne sont pas affichees aux utilisateurs normaux par le catalogue dynamique.")


def architecture(doc):
    add_heading(doc, "2. Architecture generale")
    add_heading(doc, "Pipeline logique", 2)
    add_code(doc, """Message utilisateur
  -> Detection conversation / aide
  -> Precontrole d'answerability
  -> Resolution du follow-up
  -> Classification semantic / knowledge / conversation
  -> Routage agent
  -> Extraction d'intention et resolution des entites
  -> Application du scope utilisateur et du RLS
  -> Selection d'un template DAX gouverne
  -> Execution Power BI / Power Automate ou recherche documentaire
  -> Validation du resultat et garde-fou de grounding
  -> Selection du template de reponse
  -> Persistance du message, du contexte et des artefacts
  -> Rendu interactif et actions contextuelles""")
    add_heading(doc, "Couches", 2)
    add_table(doc, ["Couche", "Responsabilite", "Composants principaux"], [
        ("Interface", "Saisie, voix, messages, chips, tables, graphiques, actions", "ai.js, templates, styles"),
        ("Conversation", "Salutations, aide, langue, follow-up, references", "conversation_intent_service, conversation_follow_up_resolution_service"),
        ("Routage", "Choix semantic/knowledge et agent", "chat_routing_service, agent_router_service"),
        ("Orchestration", "Plan d'execution et enveloppe de reponse", "powerbi_interaction_orchestrator, ai_agent_execution_service"),
        ("Donnees", "DAX gouverne, Power Automate, semantic models", "dax_generator_service, power_automate, powerbi"),
        ("Metier", "Fleet, Performance, Parts, Downtime", "services specialises"),
        ("Confiance", "Answerability, grounding, abstention", "answerability_assessment_service, grounded_response_guard_service"),
        ("Persistance", "Conversations, messages, contextes, artefacts, logs", "modeles AIConversation* et logs"),
    ], widths=[2.5, 6.2, 8], font_size=7.2)
    add_heading(doc, "Deterministe d'abord", 2)
    add_p(doc, "Les demandes evidentes de flotte, serie, equipement, KPI, periode, aide et salutations sont traitees avec des regles et des mappings. L'IA externe n'est pas necessaire pour construire les lignes d'une flotte, calculer un KPI semantique, filtrer un tableau ou generer un fichier Excel.")
    add_p(doc, "Le LLM intervient surtout pour certaines classifications, reformulations, analyses de commentaires ou reponses combinees. Une reponse deterministe reste disponible lorsque les donnees Power BI sont valides mais que le fournisseur d'IA echoue.")


def user_experience(doc):
    add_heading(doc, "3. Experience utilisateur")
    add_heading(doc, "Types d'entree", 2)
    add_bullets(doc, [
        "Question libre en francais ou en anglais.",
        "Numero de serie seul, par exemple L7K00442.",
        "Code equipement, par exemple EX007 ou HT005.",
        "Follow-up court : Only the 777, What about Essakane?, For June 2026, Download it.",
        "Commande de navigation : Open the Benchmark page ou Open saved context in Power BI.",
        "Entree vocale via l'API de transcription lorsque la fonction est disponible.",
    ])
    add_heading(doc, "Types de sortie", 2)
    add_bullets(doc, [
        "Texte naturel avec chiffres formates selon la langue.",
        "Cartes KPI et grilles multi-KPI.",
        "Tables interactives avec recherche, tri, filtre et pagination.",
        "Tendances, comparaisons, classements et Pareto.",
        "Artefacts persistants reaffichables apres reouverture de la conversation.",
        "Actions : comparer, afficher la tendance, voir les drivers, ouvrir Power BI, exporter Excel, retry.",
        "Messages d'abstention differencies : information absente, valeur vide, entite introuvable, acces refuse ou source indisponible.",
    ])
    add_heading(doc, "Langue et format", 2)
    add_p(doc, "La langue de reponse suit le dernier message. Les noms de MineSite, codes machines, modeles, numeros de serie et acronymes KPI ne sont pas traduits. Les pourcentages utilisent la convention locale, par exemple 84,25 % en francais et 84.25% en anglais.")


def agents_and_routing(doc):
    add_heading(doc, "4. Agents et routage")
    agents = AIAgent.objects.order_by("code")
    add_table(doc, ["Code", "Nom", "Type", "Actif", "Validation"], [
        (a.code, a.name, a.agent_type, bool_text(a.active), a.validation_status) for a in agents
    ], widths=[3.2, 4, 3.2, 1.5, 2.5], font_size=7.5)
    add_heading(doc, "Machine Performance Agent", 2)
    add_p(doc, "Agent principal pour les valeurs operationnelles. Il couvre Fleet Inventory, Fleet Performance, Availability, Downtime, equipements, diagnostic, navigation Power BI et export. Les requetes passent par des outils gouvernes et par les droits de l'utilisateur.")
    add_heading(doc, "Mining Knowledge Agent", 2)
    add_p(doc, "Agent documentaire pour les definitions, procedures, recommandations et Best Practices. Il recherche uniquement des contenus indexes et valides, avec references de source. Sans preuve documentaire suffisante, il s'abstient.")
    add_heading(doc, "Execution combinee", 2)
    add_p(doc, "Une question qui demande a la fois un resultat operationnel et une recommandation peut declencher Machine Performance puis Mining Knowledge. La reponse separe les constats operationnels, les recommandations documentees et l'analyse suivante proposee.")
    add_heading(doc, "Signaux de routage", 2)
    add_table(doc, ["Signal", "Route habituelle"], [
        ("KPI + filtre/periode/action", "Machine Performance / semantic_query"),
        ("Definition sans filtre operationnel", "Knowledge question"),
        ("Best Practice, procedure, according to", "Mining Knowledge"),
        ("Valeur + recommandation", "Combined"),
        ("Salutation, merci, aide", "Conversation deterministe"),
        ("Question incomplete ou concept ambigu", "Clarification"),
        ("Contexte recent compatible", "Follow-up Machine Performance"),
    ], widths=[7, 9], font_size=8)


def intent_and_parameters(doc):
    add_heading(doc, "5. Modele d'intention et parametres")
    add_p(doc, "Une demande analytique est normalisee en une structure composee. Cette structure est la frontiere entre le langage naturel et l'execution gouvernee.")
    add_code(doc, """{
  "section": "performance",
  "domain": "machine_performance",
  "capability": "fleet_performance",
  "intent_type": "single_kpi",
  "metric": "availability",
  "metrics": ["availability"],
  "metric_bundle": null,
  "scope_type": "serial_number",
  "filters": {
    "minesite": "Fekola",
    "model": "6020",
    "serial_number": "DNR00153",
    "period": "year to date"
  },
  "group_by": [],
  "comparison": {},
  "navigation": {},
  "export": false
}""")
    add_heading(doc, "Parametres d'entree de l'orchestrateur", 2)
    add_table(doc, ["Parametre", "Type", "Usage"], [
        ("question_text", "string", "Question originale; langue, intent et reponse."),
        ("user_context.user", "User", "Permissions, MineSites, RLS, feature flags."),
        ("user_context.section_code", "string", "Section AI Config cible."),
        ("user_context.dataset_name", "string", "Semantic model demande."),
        ("user_context.debug_mode", "boolean", "Trace admin detaillee."),
        ("user_context.open_report", "boolean", "Autorise la preparation de navigation."),
        ("user_context.pre_extracted_intent", "object", "Intent deja resolu par le follow-up."),
        ("user_context.follow_up_resolution", "object", "Operations et contexte herite."),
        ("conversation_context.conversation_id", "UUID/string", "Isolation et persistance."),
        ("conversation_context.messages", "array", "Historique utile au rendu ou au fournisseur."),
        ("conversation_context.validated_intent", "object", "Dernier intent valide compatible."),
    ], widths=[5, 2.8, 9], font_size=7.2)
    add_heading(doc, "Catalogue d'intents Machine Performance", 2)
    intents = sorted(INTENT_TYPES)
    rows = [(intents[i], intents[i + 1] if i + 1 < len(intents) else "") for i in range(0, len(intents), 2)]
    add_table(doc, ["Intent", "Intent"], rows, widths=[8, 8], font_size=7.5)


def entity_and_time(doc):
    add_heading(doc, "6. Resolution des entites, synonymes et periodes")
    add_heading(doc, "Entites", 2)
    add_bullets(doc, [
        "MineSite et Customer, avec valeurs autorisees et synonymes valides.",
        "Model conserve comme texte : 777, 777 WT, 307.5, PC3000-6.",
        "Equipment et Serial Number avec exact, casse, espaces puis normalisation alphanumerique unique.",
        "Equipment Family via ParentProductGroup et groupes Prime Movers via ModelList_MiningProd[PrimeMovers].",
        "KPI, periods, downtime driver, component et autres dimensions configurees.",
    ])
    add_heading(doc, "Synonymes actifs", 2)
    counts = KnowledgeSynonym.objects.filter(is_active=True).values(
        "section__code", "entity_type", "validation_status"
    ).annotate(count=Count("id")).order_by("section__code", "entity_type", "validation_status")
    add_table(doc, ["Section", "Type", "Validation", "Nombre"], [
        (x["section__code"], x["entity_type"], x["validation_status"], x["count"]) for x in counts
    ], widths=[3, 5, 3, 2], font_size=7.5)
    add_p(doc, "Les suggestions floues ne doivent pas reveler de sites ou entites hors du perimetre autorise. Une correspondance ambiguë conduit a une clarification.")
    add_heading(doc, "Periodes reconnues", 2)
    add_table(doc, ["Expression", "Valeur canonique / comportement"], [
        ("YTD, Year to Date, annee en cours", "year to date; du 1er janvier a la derniere date valide"),
        ("MTD, Month to Date", "month to date"),
        ("Current Month / ce mois", "mois courant"),
        ("Previous Month / mois precedent", "mois calendaire precedent"),
        ("June 2026 / juin 2026", "2026-06"),
        ("2026-06-15", "jour specifique"),
        ("2026", "annee specifique"),
        ("Last 12 Months", "fenetre glissante de 12 mois"),
    ], widths=[6, 10], font_size=8)
    add_callout(doc, "Regle YTD", "Le filtre est aligne sur la date disponible du modele et ne doit pas inclure un horizon futur du calendrier. La forme sauvegardee YTD 2026 est recanonisee avant generation DAX.")


def fleet_inventory(doc):
    add_heading(doc, "7. Fleet Inventory")
    add_p(doc, "Fleet Inventory signifie l'inventaire des equipements affectes a un MineSite. Il ne signifie ni Availability, ni MTBF, ni downtime. Une question comme « C'est quoi la flotte de Fekola ? » est donc routee vers les donnees de parc.")
    add_heading(doc, "Source et mapping", 2)
    add_p(doc, "Source de verite : semantic model FPR Global DB + RLS, table EquipmentList_MiningProd.")
    add_table(doc, ["Champ canonique", "Colonne source", "Affichage"], [
        ("site", "Site", "Site"), ("equipment", "Equipment", "Equipment"),
        ("model", "Model", "Model"), ("serial_number", "SN", "Serial Number"),
        ("equipment_family", "ParentProductGroup", "Equipment Family"),
        ("brand", "Brand", "Brand"), ("equipment_id", "EquipID", "Equipment ID"),
        ("status", "Status", "Status not mapped tant qu'aucune regle n'est validee"),
        ("smu", "SMU.SMU", "SMU; valeur vide conservee comme indisponible"),
    ], widths=[4, 5, 7], font_size=7.5)
    add_heading(doc, "Questions prises en charge", 2)
    add_bullets(doc, [
        "Flotte complete d'un site.", "Flotte regroupee par modele.",
        "Equipements d'un modele exact sur un site.", "Nombre de machines et de modeles.",
        "Recherche par serial number, y compris un serial saisi seul.",
        "Recherche par code equipement.", "Export Excel du resultat sauvegarde.",
    ])
    add_heading(doc, "Reponse et deduplication", 2)
    add_p(doc, "La table utilisateur impose Site, Equipment, Model, Serial Number dans cet ordre. Les champs optionnels apparaissent dans la fiche machine. L'identite de deduplication privilegie EquipID, puis un serial unique normalise, puis la cle composite Site + Equipment + Model + Serial Number.")
    add_p(doc, "La limite de securite est 5 000 lignes par defaut, plafonnee techniquement entre 100 et 50 000. Le service recupere une ligne supplementaire pour detecter une troncature et refuse de presenter une liste incomplete comme complete.")
    add_heading(doc, "Valeurs manquantes", 2)
    add_table(doc, ["Champ", "Rendu"], [
        ("Serial Number", "Not available"), ("Model", "Unknown Model"),
        ("Equipment", "Not available"), ("Family", "Not classified"),
        ("Brand", "Not available"), ("SMU", "Not available, jamais zero par defaut"),
        ("Status", "Status not mapped"),
    ], widths=[5, 11], font_size=8)


def fleet_performance(doc):
    add_heading(doc, "8. Fleet Performance Intelligence")
    add_p(doc, "Le moteur combine une operation analytique, un ou plusieurs KPI, un scope, des filtres, un groupement et une periode. Les mesures de ratio sont evaluees directement dans le contexte Power BI; elles ne sont pas moyennees manuellement en Python.")
    add_heading(doc, "Bundle principal", 2)
    add_table(doc, ["Code", "Libelle", "Mesure Power BI", "Unite", "Precision", "Direction runtime"], [
        (code, item["label"], item["measure"], item["unit"], item["precision"], item["optimization"])
        for code, item in CORE_METRICS.items()
    ], widths=[3.2, 3.5, 4.2, 2, 1.5, 2.5], font_size=7)
    add_heading(doc, "Bundles", 2)
    add_table(doc, ["Bundle", "KPI"], [(code, list_text(metrics)) for code, metrics in METRIC_BUNDLES.items()], widths=[5, 11], font_size=8)
    add_heading(doc, "Operations analytiques", 2)
    add_bullets(doc, [
        "Vue generale six KPI pour une flotte, un site, un modele ou une machine.",
        "Valeur d'un KPI unique avec le KPI demande comme sortie principale.",
        "Reliability overview : MTBS, MTBF et MTTR.",
        "Planned versus Unplanned sans normalisation artificielle a 100 %.",
        "Comparaison entre sites, modeles, equipements, serials ou periodes.",
        "Tendance mensuelle d'un ou plusieurs KPI.",
        "Classement ascendant ou descendant, avec exclusion des valeurs vides.",
        "Benchmark contre des pairs autorises; aucun benchmark industriel invente.",
        "SMU tracking comme valeur master data, sans interpretation automatique en utilisation.",
    ])
    add_heading(doc, "Couverture", 2)
    add_p(doc, "Le payload peut distinguer les equipements de la flotte master, les equipements avec donnees de performance, ceux sans donnees et le pourcentage de couverture. Une absence de performance n'est pas convertie en zero.")
    add_heading(doc, "Templates DAX", 2)
    templates = AIDaxTemplate.objects.filter(section__code="performance", is_active=True).order_by("template_code")
    add_table(doc, ["Code", "Nom", "Usage"], [
        (t.template_code, t.template_name, t.description) for t in templates if t.template_code.startswith(("PERF_", "BUNDLE_"))
    ], widths=[5, 4, 8], font_size=6.8)


def availability(doc):
    add_heading(doc, "9. Physical Availability")
    add_p(doc, "La mesure officielle est [Availability New]. Elle est evaluee dans le contexte Site, Model, Equipment, Serial Number, Family/Product Group et periode. Le rendu est un pourcentage a deux decimales.")
    add_heading(doc, "Scopes", 2)
    add_bullets(doc, [
        "Flotte globale autorisee.", "MineSite ou Customer.", "Model ou famille.",
        "Prime Mover group, par exemple HMS.", "Equipment ou Serial Number.", "Periode et tendance." ])
    add_heading(doc, "Traitement YTD", 2)
    add_p(doc, "Pour Availability YTD, le generateur utilise une expression alignee avec le rapport : evaluation mensuelle de la mesure puis ponderation par le nombre de jours des mois disponibles. Cette regle evite de moyenner arbitrairement des valeurs machine ou des pourcentages affiches.")
    add_heading(doc, "Alias Prime Movers", 2)
    add_table(doc, ["Libelle", "Code filtre"], [
        ("Hydraulic Mining Shovels", "HMS"), ("Large Mining Trucks", "LMT"),
        ("Large Track Type Tractors", "LTTT"), ("Large Wheel Loaders", "LWL"),
        ("Motor Graders", "MG"), ("Off Highway Trucks", "OHT"),
        ("Surface Rotary Drills", "SRD"), ("Wheel Dozers", "WD"),
        ("Excavators", "EXC"),
    ], widths=[9, 4], font_size=8)
    add_callout(doc, "Correction importante", "Hydraulic Mining Shovels est un libelle utilisateur. Le filtre semantique valide est ModelList_MiningProd[PrimeMovers] = HMS, et non EquipmentList_MiningProd[ParentProductGroup] = Hydraulic Mining Shovels.")


def downtime(doc):
    add_heading(doc, "10. Downtime, Pareto et Root Cause")
    add_p(doc, "Le diagnostic downtime reutilise le contexte analytique courant. Il execute la mesure [DonwtimeHours] sur DowntimeData_MiningProd et groupe par DescriptionCat. Il retourne les heures, occurrences, duree moyenne, equipements affectes, part et cumul Pareto.")
    add_heading(doc, "Filtres et sorties", 2)
    add_table(doc, ["Element", "Detail"], [
        ("Work Type", "Planned ou Unplanned; valeur injectee par TREATAS sur DowntimeData_MiningProd[WorkType]."),
        ("Top N", "10 par defaut, contraint de 1 a 25 dans le diagnostic Availability."),
        ("Driver", "DescriptionCat; les lignes vides et les heures <= 0 sont exclues."),
        ("Pareto", "Share of downtime et cumulative percentage calcules sur le total semantique."),
        ("Affected Equipment", "Distinct count de SN."),
        ("Events", "Liste gouvernee limitee, avec commentaires lorsque disponibles."),
    ], widths=[4, 12], font_size=7.8)
    add_heading(doc, "Root Cause Explorer", 2)
    add_bullets(doc, [
        "Resume de la session et selections actives.", "Breakdown par dimensions configurees.",
        "Equipements affectes et liste d'evenements.", "Analyse des commentaires avec couverture.",
        "Detection de pannes repetees.", "Classification SMCS explicite ou resolue.",
        "Navigation, retour et reset de session." ])
    add_p(doc, "Une cause probable ou un theme de commentaire ne doit pas etre presente comme une cause racine confirmee sans preuve et validation metier.")


def parts_sales(doc):
    add_heading(doc, "11. Parts Sales")
    config = BusinessPerformanceConfig.objects.filter(is_active=True).first()
    add_p(doc, "Le chatbot Parts Sales interroge le semantic model Mine Logistics & AfterMarket via BusinessPerformanceService. La requete applique le domaine LOB Parts et les canaux de distribution configures.")
    if config:
        add_table(doc, ["Parametre", "Valeur active"], [
            ("Semantic model", config.semantic_model_name),
            ("Semantic model ID", config.semantic_model_id),
            ("Authentification", config.authentication_mode),
            ("Devise par defaut", config.default_currency),
            ("LOB Parts", config.parts_lob_values),
            ("Canaux directs", config.direct_sales_channel_values),
            ("Cache", f"{config.cache_duration_seconds} s"),
            ("Timeout requete", f"{config.query_timeout_seconds} s"),
        ], widths=[5, 11], font_size=8)
    add_heading(doc, "Mesures de facture", 2)
    mappings = BusinessPerformanceMapping.objects.filter(
        logical_name__in=["global_revenue_eur", "global_revenue_usd", "global_revenue_cfa"]
    ).order_by("display_order")
    add_table(doc, ["Code", "Mesure", "Format", "Active"], [
        (m.logical_name, m.object_name, m.format_string, bool_text(m.is_active)) for m in mappings
    ], widths=[5, 5, 3, 2], font_size=8)
    add_p(doc, "Le service conversationnel Parts Sales actuel force currency=EURO et affiche CA Facture EU. Les mappings USD et CFA existent au niveau Business Performance, mais leur selection conversationnelle doit etre confirmee avant de les annoncer comme disponibles dans toutes les formulations.")
    add_heading(doc, "Scopes", 2)
    add_bullets(doc, [
        "Total Parts YTD.", "Parts YTD par customer.", "Parts YTD par territoire/MineSite.",
        "Fallback customer lorsqu'un site est represente par le nom client.",
        "Top 10 lorsque l'utilisateur demande une ventilation sans valeur precise." ])


def knowledge_and_reporting(doc):
    add_heading(doc, "12. Knowledge, Best Practices et Reporting")
    add_heading(doc, "Knowledge Base", 2)
    add_p(doc, "La recherche documentaire utilise les ressources validees, un score lexical et, selon configuration, des embeddings. Les resultats peuvent fournir recommandations, extraits, titre et page. En l'absence de source validee, le chatbot renvoie INSUFFICIENT_EVIDENCE au lieu de repondre depuis sa memoire generale.")
    add_bullets(doc, [
        "Definitions KPI et glossaire metier.", "Terminologie miniere.",
        "Best Practices et recommandations documentees.", "Procedures et regles metier.",
        "Citations des documents et pages lorsque disponibles.", "Reponse combinee avec constat operationnel." ])
    add_heading(doc, "Reporting et Power BI", 2)
    add_p(doc, "Le chatbot peut rechercher un rapport configure, preparer une navigation et ouvrir le viewer generique avec un contexte valide. La navigation repose sur les identifiants internes synchronises, les mappings de pages/visuels et les permissions du rapport.")
    add_bullets(doc, [
        "Recherche et selection de rapport.", "Ouverture du rapport et de la page configuree.",
        "Application de filtres valides au viewer.", "Conservation du contexte conversationnel.",
        "Lien Open saved context in Power BI.", "Pas de DAX arbitraire genere par le LLM." ])


def capability_answerability(doc):
    add_heading(doc, "13. Aide, capability discovery et answerability")
    add_heading(doc, "Capability Discovery", 2)
    add_p(doc, "Les questions « Que peux-tu faire ? », « Help » ou « What can you do with this machine? » sont resolues sans requete Power BI et sans appel OpenAI. Le catalogue est charge depuis AIAgentCapability puis filtre selon l'agent, les permissions, les feature flags, la readiness et le contexte courant.")
    caps = AIAgentCapability.objects.select_related("agent").filter(
        readiness_status__in=["Ready", "Limited"], enabled=True
    ).order_by("display_order")
    add_table(doc, ["Code", "Categorie", "Etat", "Entites", "Operations"], [
        (c.capability_code, c.category, c.readiness_status, list_text(c.supported_entities_json), list_text(c.supported_operations_json)) for c in caps
    ], widths=[3.2, 3.5, 2, 3, 5], font_size=6.8)
    add_heading(doc, "Statuts d'answerability", 2)
    statuses = [
        ("ANSWERABLE", "Information recuperee et suffisamment etayee."),
        ("NEEDS_CLARIFICATION", "Entite ou parametre necessaire manquant/ambigu."),
        ("INFORMATION_NOT_IN_CONFIGURED_SOURCES", "Champ absent des sources actives."),
        ("FIELD_AVAILABLE_BUT_VALUE_MISSING", "Champ configure mais valeur vide pour l'entite."),
        ("ENTITY_NOT_FOUND", "Aucune entite autorisee correspondante."),
        ("ACCESS_RESTRICTED", "Utilisateur non autorise; aucune fuite de detail."),
        ("CAPABILITY_NOT_CONFIGURED", "Fonction conceptuelle mais mappings incomplets."),
        ("TEMPORARILY_UNAVAILABLE", "Source/service temporairement indisponible; retry pertinent."),
        ("INSUFFICIENT_EVIDENCE", "Preuve structuree ou documentaire insuffisante."),
        ("CONFLICTING_SOURCES", "Sources incompatibles sans priorite validee."),
        ("UNSUPPORTED_ACTION", "Action non supportee."),
        ("OUT_OF_SCOPE", "Question hors des domaines actifs."),
        ("LOW_CONFIDENCE", "Resolution sous le seuil de confiance."),
    ]
    add_table(doc, ["Statut", "Signification"], statuses, widths=[6, 10], font_size=7.4)
    config = AIAnswerabilityConfiguration.objects.filter(active=True).first()
    if config:
        add_heading(doc, "Parametres actifs", 2)
        add_table(doc, ["Parametre", "Valeur"], [
            ("minimum_entity_confidence", config.minimum_entity_confidence),
            ("minimum_knowledge_confidence", config.minimum_knowledge_confidence),
            ("require_structured_evidence", bool_text(config.require_structured_evidence)),
            ("require_document_evidence", bool_text(config.require_document_evidence)),
            ("allow_general_model_knowledge", bool_text(config.allow_general_model_knowledge)),
            ("allow_hypothesis_mode", bool_text(config.allow_hypothesis_mode)),
            ("show_available_alternatives", bool_text(config.show_available_alternatives)),
            ("show_source_summary", bool_text(config.show_source_summary)),
            ("log_data_gaps", bool_text(config.log_data_gaps)),
            ("allow_data_gap_reporting", bool_text(config.allow_data_gap_reporting)),
        ], widths=[7, 5], font_size=8)
    add_heading(doc, "Exemple Commissioning Date", 2)
    add_p(doc, "Le champ commissioning_date existe dans le registre avec l'etat Not Configured et source_type=none. Si DT677 existe, le chatbot doit indiquer que la date de mise en service n'est pas disponible dans les sources configurees. Il ne doit pas utiliser une premiere panne, une premiere valeur SMU ou une date de refresh comme substitut.")
    add_heading(doc, "Grounded Response Guard", 2)
    add_p(doc, "Le garde-fou recherche dans la reponse proposee les dates, nombres et identifiants. Chaque fait doit apparaitre dans l'evidence structuree. Une affirmation non supportee est remplacee par un message deterministic d'insuffisance de preuve et journalisee.")


def persistence_exports(doc):
    add_heading(doc, "14. Conversations, contextes, artefacts et exports")
    add_heading(doc, "Persistance", 2)
    add_bullets(doc, [
        "AIConversation : proprietaire, statut, titre, contexte actif et dates.",
        "AIConversationMessage : role, contenu, statut, agent et metadata.",
        "AIConversationContext : intent valide actif, filtre et agent courant.",
        "AIConversationArtifact : snapshot immutable du resultat et version de refresh.",
        "Archive/restore de conversation et retry de message.",
        "Les resultats historiques ne sont pas rerun automatiquement lors de la reouverture." ])
    add_heading(doc, "Follow-up", 2)
    add_p(doc, "Le resolver recupere le dernier contexte compatible reussi et applique des operations set, append, clear, keep ou compare_with. Il peut remplacer uniquement le site, modele, periode ou KPI, conserver le reste du contexte et exporter le dernier artefact compatible.")
    add_table(doc, ["Follow-up", "Operation"], [
        ("Only the 777", "set_model=777"), ("What about Essakane?", "set_site=Essakane"),
        ("For June 2026", "set period=2026-06"), ("All models", "clear_model"),
        ("Compare with Fekola", "entity comparison"), ("Download it", "export current compatible artifact"),
        ("Same model", "inherit selected equipment model"), ("Show its drivers", "reuse active equipment and switch intent"),
    ], widths=[6, 10], font_size=8)
    add_heading(doc, "Fleet Excel", 2)
    add_p(doc, "L'export Fleet est reconstruit depuis l'artefact fleet_equipment_table appartenant a l'utilisateur. Il contient Fleet by Model, Fleet Details et Export Metadata. Model et Serial Number sont forces en texte, les entetes sont gelees et les filtres Excel sont actifs. Generated By vaut Mining360.")
    add_heading(doc, "Fleet Performance Excel", 2)
    add_p(doc, "L'export Performance utilise l'artefact fleet_performance_analysis et cree Performance Summary, Performance Results et Export Metadata. Il preserve les lignes et metriques du resultat sauvegarde. Aucun nouveau DAX moins contraint n'est execute pour l'export.")
    add_heading(doc, "Securite des exports", 2)
    add_bullets(doc, [
        "Authentification obligatoire.", "Ownership conversation/artefact controle.",
        "Feature flag et permission reevalues au moment du download.", "Identifiant d'export opaque.",
        "Nom de fichier sanitise.", "Aucun token, secret ou chemin serveur dans la reponse." ])


def security(doc):
    add_heading(doc, "15. Autorisations, scope MineSite et RLS")
    add_p(doc, "Le perimetre est applique avant l'execution. Un utilisateur ayant le role business_performance_role=MineSite doit posseder exactement un MineSite dans business_performance_scope et une identite effective valide.")
    add_numbered(doc, [
        "Charger le profil PlatformUser et determiner si l'utilisateur est admin ou restreint.",
        "Comparer le MineSite demande aux sites autorises, sans suggerer de site interdit.",
        "Forcer le MineSite autorise dans l'intent lorsqu'il est restreint.",
        "Resoudre les roles RLS du semantic model.",
        "Transmettre roles et effectiveUser au flux Power BI.",
        "Filtrer les capacites, rapports, exports et artefacts selon le meme profil.",
    ])
    add_callout(doc, "Non-divulgation", "Une reponse ACCESS_RESTRICTED ne confirme pas que l'equipement existe sur un autre site. Les autocomplete, fuzzy matches, nombres et exports doivent rester dans le scope autorise.", RED)
    add_heading(doc, "Modes Power BI", 2)
    add_table(doc, ["Mode", "Token", "Identite"], [
        ("App Owns Data", "Embed Token", "Service principal + effective identity/RLS selon configuration"),
        ("User Owns Data", "Delegated Entra token", "Utilisateur connecte; refresh delegue"),
        ("Power Automate query", "Endpoint securise configure", "Dataset ID, roles et effective user transmis"),
    ], widths=[4, 4, 8], font_size=8)




def config_admin(doc):
    add_heading(doc, "17. AI Config et gouvernance")
    add_p(doc, "AI Config centralise les metadonnees qui transforment une question en execution controlee. Les administrateurs peuvent maintenir les sections, mappings, exemples, synonymes, templates, agents, capabilities et regles de confiance.")
    add_table(doc, ["Repository", "Usage"], [
        ("AIConfigSection", "Domaines Performance, Parts Sales, Planned Component Rebuild, Power BI Reporting."),
        ("AIMetricMapping", "Code metrique -> mesure Power BI."),
        ("AIFilterMapping", "Code filtre -> table, colonne, type."),
        ("AIDaxTemplate", "Registre de templates fermes et gouvernes."),
        ("AIResponseTemplate", "Composition de l'artefact et ordre des composants."),
        ("AIIntentResponseTemplateMapping", "Intent/scope/metrique -> template de reponse."),
        ("KnowledgeKPIDictionary", "Definition, mesure, unite, precision, agrégation, dimensions et seuils."),
        ("KnowledgeSynonym", "Valeurs metier, variantes linguistiques et statut de validation."),
        ("BusinessDataField", "Couverture des champs et source de verite."),
        ("AIAgentCapability", "Readiness, permissions, feature flags, exemples et operations."),
        ("AIAnswerabilityConfiguration", "Seuils et politiques de preuve."),
        ("UnansweredInformationRequirement", "Registre agrege des besoins de donnees manquantes."),
    ], widths=[5, 11], font_size=7.5)


def production_hardening(doc):
    add_heading(doc, "18. Durcissement de production du chatbot")
    add_p(doc, "L'interface ne presente plus une fonction parce qu'un intent ou un template existe. Une suggestion visible doit etre rattachee a une operation prete, une configuration valide, des permissions suffisantes, des dependances disponibles et une certification de bout en bout encore valide.")
    add_heading(doc, "Quatre niveaux de readiness", 2)
    add_table(doc, ["Niveau", "Objet", "Decision utilisateur"], [
        ("Capability", "Famille large : Fleet, Performance, Downtime, Knowledge, Reporting.", "Peut etre Ready ou Limited sans rendre toutes ses operations visibles."),
        ("Operation", "Action analytique precise et ses dependances.", "READY est requis pour une suggestion directe de production."),
        ("Suggestion", "Libelle, type d'action, contexte, contrat attendu et certification.", "CERTIFIED est requis dans l'environnement courant."),
        ("Action", "Bouton contextuel rattache a un artefact et une route.", "Rendue seulement apres ActionEligibilityService."),
    ], widths=[2.5, 6.5, 7], font_size=7.2)
    add_heading(doc, "Registres gouvernes", 2)
    add_bullets(doc, [
        "AICapabilityOperation decrit les entites, metriques, sources, templates, flags, permissions, artefacts et fallback requis.",
        "AIChatSuggestion remplace les prompts statiques et definit DIRECT, GUIDED, CONTEXTUAL, NAVIGATION ou TOOL.",
        "AISuggestionCertification conserve l'environnement, les versions, les attentes, les resultats HTTP, grounding, persistance, UI et permissions.",
        "AIActionContract gouverne les boutons contextuels.",
        "FeaturePilotMembership implemente les vrais utilisateurs et groupes Pilot avec dates d'effet.",
        "AIDependencyHealthSnapshot fournit un etat cache sans appel distant au rendu.",
        "AIConversationExecution et AIChatInteractionEvent assurent machine d'etats, audit, latence et supervision." ])
    operations = AICapabilityOperation.objects.select_related("capability").order_by("domain_code", "operation_code")
    add_heading(doc, "Readiness des operations observees", 2)
    add_table(doc, ["Capability", "Operation", "Readiness", "Score", "Contexte zero", "Validation"], [
        (item.capability.capability_code if item.capability else item.domain_code, item.operation_code, item.readiness_status, item.readiness_score, "Oui" if item.supports_zero_context else "Non", item.validation_status)
        for item in operations
    ], widths=[3, 4, 3, 1.5, 2, 2.5], font_size=6.8)
    suggestions = AIChatSuggestion.objects.select_related("operation").order_by("display_order", "suggestion_code")
    add_heading(doc, "Suggestions migrees et eligibilite", 2)
    add_table(doc, ["Code", "Libelle EN", "Type", "Operation", "Tier", "Certification", "Visible"], [
        (item.suggestion_code, item.label_en, item.action_type, item.operation.operation_code, item.reliability_tier, item.certification_status,
         "Oui, si scope et flags" if item.certification_status == "CERTIFIED" and item.operation.readiness_status == "READY" else "Non")
        for item in suggestions
    ], widths=[4, 4, 2.2, 3.2, 1, 2.3, 2], font_size=6.2)
    add_callout(doc, "Analyze repeated failures", "La suggestion est FAILED et l'operation repeated_failures est Needs Configuration. Elle est retiree du nouveau chat. Une question manuelle recoit une abstention CAPABILITY_NOT_CONFIGURED sans lancer Power BI. Elle ne pourra etre publiee qu'apres contexte guide, contrat d'evidence, fallback et certification E2E.")
    add_heading(doc, "Cycle d'execution controle", 2)
    add_code(doc, """Suggestion eligible
  -> idempotency_key + client_execution_id
  -> message utilisateur persiste
  -> placeholder assistant persiste
  -> ROUTING -> CHECKING_ANSWERABILITY
  -> source ou registre local
  -> VALIDATING_GROUNDING
  -> SUCCEEDED / NEEDS_CLARIFICATION / ABSTAINED / RETRYABLE_FAILED
  -> actions filtrees
  -> evenement et latence
  -> auto-invalidation si le taux de succes chute""")
    add_p(doc, "Un double clic ou une repetition avec la meme cle retourne l'execution existante. Une soumission sans identifiant reutilise la derniere conversation active vide avant de consommer un nouveau slot. La limite de conversations est une issue explicite non retryable, et non une panne de source. Cancel conserve le message utilisateur et empeche l'insertion tardive d'un resultat stale. Retry est reserve aux pannes transitoires et cree une nouvelle execution liee a l'ancienne.")
    add_heading(doc, "Etats utilisateur et robustesse UI", 2)
    add_bullets(doc, [
        "Les suggestions sont chargees depuis /api/ai/chat/suggestions/ sans Power BI, Knowledge Search ni fournisseur IA.",
        "Les suggestions guidees affichent uniquement des valeurs d'entite autorisees.",
        "Le clic utilise le meme endpoint et la meme persistence que la saisie manuelle.",
        "Le chargement est explicite, puis propose Keep waiting ou Cancel apres delai.",
        "Les erreurs techniques sont remplacees par des issues controlees et localisees.",
        "Le responsive mobile conserve le chatbot dans le viewport et le composer hors du contenu.",
        "L'auto-masquage invalide une suggestion certifiee apres au moins trois outcomes recents sous le seuil de succes configure." ])


def feature_flag_section(doc):
    add_heading(doc, "19. Feature flags et rollout")
    add_p(doc, "Les valeurs Production activent la fonction pour les utilisateurs autorises. Admin Only la limite aux administrateurs de plateforme. Pilot inclut les administrateurs ainsi que les utilisateurs ou groupes actifs declares dans FeaturePilotMembership.")
    add_table(doc, ["Feature flag", "Valeur active"], feature_flags(), widths=[11, 4], font_size=7)
    add_callout(doc, "Impact actuel", "Fleet Inventory, complete Fleet Performance, exports, Capability Discovery et Answerability restent Admin Only ou Pilot dans la configuration par defaut. Ils ne doivent pas etre annonces comme disponibles a tous les utilisateurs tant que le rollout n'est pas passe a Production.")


def api_contracts(doc):
    add_heading(doc, "20. API et contrats de reponse")
    add_heading(doc, "Endpoints principaux", 2)
    endpoints = [
        ("POST /ai/ask/", "Question principale et routage."),
        ("GET /api/ai/chat/suggestions/", "Suggestions certifiees, localisees et autorisees."),
        ("POST /api/ai/chat/suggestions/events/", "Telemetrie UI controlee."),
        ("GET /api/ai/chat/readiness/", "Snapshot administrateur de readiness."),
        ("POST /api/ai/chat/executions/{id}/cancel/", "Annulation d'une execution appartenant a l'utilisateur."),
        ("GET/POST /api/ai/conversations/", "Liste et creation de conversations."),
        ("GET /api/ai/conversations/{id}/messages/", "Historique des messages."),
        ("GET /api/ai/conversations/{id}/artifacts/", "Artefacts persistants."),
        ("POST /api/ai/conversations/{id}/messages/{id}/retry/", "Retry d'un echec transitoire."),
        ("POST /api/fleet/exports/", "Creation export Fleet depuis artefact."),
        ("GET /api/fleet/exports/{id}/download/", "Telechargement Fleet securise."),
        ("POST /api/performance/exports/", "Creation export Performance."),
        ("GET /api/performance/exports/{id}/download/", "Telechargement Performance."),
        ("POST /api/ai/data-gaps/report/", "Signalement d'un besoin de donnee."),
        ("POST /api/ai/audio/transcribe/", "Transcription vocale."),
        ("POST /ai/downtime-explorer/open/", "Ouverture d'une session Root Cause."),
        ("GET /ai/downtime-explorer/{id}/summary/", "Resume Root Cause."),
        ("GET /ai/downtime-explorer/{id}/breakdown/", "Ventilation dimensionnelle."),
        ("GET /ai/downtime-explorer/{id}/equipment/", "Equipements affectes."),
        ("GET /ai/downtime-explorer/{id}/events/", "Evenements downtime."),
        ("POST /powerbi-interaction/embed-config/{report_id}/", "Configuration d'embed securisee."),
    ]
    add_table(doc, ["Endpoint", "Fonction"], endpoints, widths=[8, 8], font_size=7.2)
    add_heading(doc, "Enveloppe type", 2)
    add_code(doc, """{
  "ok": true,
  "conversation_id": "...",
  "answer": "...",
  "intent": {"intent_type": "...", "filters": {}},
  "answerability": {"status": "ANSWERABLE", "reason_code": null},
  "presentation": {"template_code": "...", "template_version": "1.0"},
  "metrics": [],
  "rows": [],
  "artifacts": [],
  "actions": [],
  "navigation": {},
  "source_summary": {},
  "warnings": [],
  "debug": {}
}""")


def errors_observability(doc):
    add_heading(doc, "21. Erreurs, reprise et observabilite")
    add_heading(doc, "Classification utilisateur", 2)
    add_table(doc, ["Situation", "Comportement"], [
        ("Champ non configure", "Information not in configured sources; pas de Retry."),
        ("Champ configure mais vide", "Field available but value missing."),
        ("Entite introuvable", "Entity not found dans le scope autorise."),
        ("Acces refuse", "Message neutre, sans confirmer l'existence."),
        ("Timeout / 5xx / rate limit", "Temporarily unavailable et Retry."),
        ("Provider IA en echec apres donnees valides", "Template deterministe; les donnees restent visibles."),
        ("Mapping manquant", "Capability not configured; detail seulement en Admin Debug."),
        ("Demande annulee ou stale", "Ignoree; pas d'erreur utilisateur."),
        ("Evidence insuffisante", "Abstention grounded."),
    ], widths=[6, 10], font_size=7.5)
    add_heading(doc, "Logs", 2)
    add_bullets(doc, [
        "AIAgentExecutionLog : agent, routage, confidence, intent, outils, sources, statut et temps.",
        "PowerBIInteractionLog : question, intent extrait/valide, DAX, dataset, statut et latence.",
        "AIAnswerabilityEvent : outcome, reason code, capability et metadata non sensible.",
        "UnansweredInformationRequirement : besoin agrege, occurrences et workflow de resolution.",
        "Artefacts : source, template, filtres, timestamp, donnees et version de refresh." ])
    add_p(doc, "Les tokens, credentials, secrets et payloads non autorises ne doivent jamais etre logs. Les erreurs provider detaillees restent reservees au debug administrateur.")


def performance_testing(doc):
    add_heading(doc, "22. Performance, cache et qualite")
    add_heading(doc, "Performance", 2)
    add_bullets(doc, [
        "Capability overview : local, sans Power BI et sans provider IA.",
        "Une seule requete detail Fleet par site lorsque possible; resume par modele calcule sur les lignes retournees.",
        "Une requete multi-mesures pour les six KPI lorsque le modele le permet.",
        "Business Performance cache actif 300 secondes dans la configuration courante.",
        "Downtime Explorer cache 600 secondes.",
        "Exports depuis artefact, sans nouvelle requete semantique.",
        "Filtrage de table local pour la recherche normale." ])
    add_heading(doc, "Tests valides", 2)
    add_p(doc, "Au 3 septembre 2026, les 487 tests de l'application reports et les 30 tests deployment sont valides. Cinq tests utilisant les fichiers temporaires Windows ont ete rejoues hors sandbox et passent. Les 86 tests cibles du chatbot passent. Les controles Django et migrations sont propres.")
    add_heading(doc, "Familles de tests", 2)
    add_bullets(doc, [
        "Routing semantic versus knowledge.", "Availability par site, modele, equipment et serial.",
        "Fleet Inventory et export Excel.", "Complete Fleet Performance et bundles KPI.",
        "Follow-up et persistance.", "Capability Discovery et commissioning date.",
        "Grounding et hallucination.", "Viewer Power BI et auth User Owns Data.",
        "Downtime Explorer, work type et Pareto.", "Permissions, scopes MineSite et non-divulgation." ])


def limitations(doc):
    add_heading(doc, "23. Limites et points de vigilance")
    rows = [
        ("Feature flags", "Plusieurs fonctions restent Admin Only/Pilot; disponibilite utilisateur non generale."),
        ("Agent validation", "Machine Performance et Mining Knowledge sont actifs mais statut To Review."),
        ("Provider", "OpenAI est actuellement marque degraded; GLM-5 est actif."),
        ("Capability registry", "Des capacites historiques sont Needs Configuration et masquees aux utilisateurs."),
        ("Suggestions Production", "Seule une suggestion certifiee dans l'environnement courant est affichee. Les autres restent masquees jusqu'a certification E2E."),
        ("Repeated failures", "Operation Needs Configuration et suggestion FAILED; absente du nouveau chat et bloquee proprement en saisie manuelle."),
        ("PM Best Practices", "Suggestion masquee tant qu'aucun corpus valide, indexe et certifie ne garantit une reponse citee."),
        ("Commissioning date", "Champ explicitement Not Configured; aucune date ne doit etre inferee."),
        ("Fleet status", "Aucune interpretation de Status n'est validee; affichage Status not mapped."),
        ("Fleet freshness", "last_refresh_at n'est pas encore alimente dans le payload Fleet Inventory."),
        ("Fleet pagination source", "Au-dela de la limite de lignes, le service refuse le resultat; pagination semantique complete reste a ajouter."),
        ("Fuzzy site", "Flag Pilot, mais la resolution Fleet principale reste surtout exacte/normalisee."),
        ("PM et Component", "Certaines intentions existent, mais le generateur peut retourner not fully configured si mesures/dimensions gouvernees manquent."),
        ("Power BI pages", "Aucune entree AIPowerBIPage active observee dans la base locale; verifier les autres mappings synchronises avant navigation fine."),
        ("Parts currency", "Mappings EURO/USD/CFA disponibles, mais le service conversationnel actuel force EURO."),
        ("Canal Parts", "Valeurs actives Onshore,Offshore; confirmer la correspondance metier avec le libelle Direct."),
        ("KPI Dictionary", "Les six KPI sont encore To Review. Des directions DB sont incoherentes avec le runtime, notamment MTTR et Unplanned; gouvernance a corriger."),
        ("Planned %", "La direction est target_based dans le runtime; aucun meilleur/moins bon sans cible validee."),
    ]
    add_table(doc, ["Sujet", "Etat / action"], rows, widths=[4, 12], font_size=7.4)
    add_callout(doc, "Important", "Le document decrit ce qui est code et configure au moment de sa generation. Toute modification AI Config, mapping Power BI, permission ou feature flag peut changer la capacite effective sans redeploiement.", RED)


def appendices(doc):
    add_heading(doc, "Annexe A. Mappings metriques actifs")
    mappings = AIMetricMapping.objects.filter(is_active=True).select_related("section").order_by("section__code", "metric_code")
    add_table(doc, ["Section", "Code", "Libelle", "Mesure Power BI"], [
        (m.section.code, m.metric_code, m.metric_label, m.powerbi_measure_name) for m in mappings
    ], widths=[3, 4, 4.5, 5], font_size=7)

    add_heading(doc, "Annexe B. Mappings filtres actifs")
    mappings = AIFilterMapping.objects.filter(is_active=True).select_related("section").order_by("section__code", "filter_code")
    add_table(doc, ["Section", "Code", "Libelle", "Table", "Colonne", "Type", "Requis"], [
        (m.section.code, m.filter_code, m.filter_label, m.powerbi_table_name, m.powerbi_column_name, m.data_type, bool_text(m.is_required)) for m in mappings
    ], widths=[2.2, 2.8, 3, 3.3, 3.5, 1.5, 1.2], font_size=6.2)

    add_heading(doc, "Annexe C. Registre des champs metier")
    fields = BusinessDataField.objects.filter(active=True).order_by("entity_type", "source_priority", "canonical_field_code")
    add_table(doc, ["Code", "Entite", "Type", "Source", "Objet", "Config.", "Validation", "Nullable"], [
        (f.canonical_field_code, f.entity_type, f.data_type, f.source_type, f"{f.table_name}.{f.column_name or f.measure_name}".strip("."), f.configuration_status, f.validation_status, bool_text(f.nullable)) for f in fields
    ], widths=[3, 2.3, 1.7, 2.5, 4, 2.2, 2, 1.3], font_size=6.2)

    add_heading(doc, "Annexe D. Templates de reponse et intents")
    templates = AIResponseTemplate.objects.filter(active=True).order_by("domain", "code")
    add_table(doc, ["Domaine", "Code", "Intent(s)", "Composant principal", "Ordre", "Version", "Validation"], [
        (t.domain, t.code, list_text(t.supported_intent_types), t.primary_component, list_text(t.component_order_json), t.version, t.validation_status) for t in templates
    ], widths=[2.7, 3.5, 4, 3, 6, 1.2, 2], font_size=5.8)

    add_heading(doc, "Annexe E. Mapping Intent -> Template")
    mappings = AIIntentResponseTemplateMapping.objects.filter(active=True).select_related("response_template").order_by("-priority", "intent_type")
    add_table(doc, ["Domaine", "Intent", "Scope", "Metrique", "Template", "Priorite", "Validation"], [
        (m.domain, m.intent_type, m.scope_type or "*", m.metric_code or "*", m.response_template.code, m.priority, m.validation_status) for m in mappings
    ], widths=[3, 4, 2, 2, 4, 1.5, 2], font_size=6)

    add_heading(doc, "Annexe F. Exemples de questions")
    examples = [
        ("Fleet Inventory", "C'est quoi la flotte de Fekola ?", "Inventaire + resume par modele + table + export."),
        ("Fleet Inventory", "Give me the Fekola 777 fleet.", "Equipements Model exact 777 a Fekola."),
        ("Machine lookup", "L7K00442", "Fiche machine si correspondance unique."),
        ("Performance", "Donne-moi la performance de Fekola YTD.", "Six KPI core."),
        ("Single KPI", "Give me the physical availability of DNR00153 on YTD.", "Availability machine et periode YTD."),
        ("Reliability", "Give me MTBS, MTBF and MTTR for Essakane 785.", "Bundle reliability_core."),
        ("Comparison", "Compare Fekola and Essakane fleet performance YTD.", "Comparaison six KPI par site."),
        ("Trend", "Show the MTBF trend for Fekola over the last 12 months.", "Tendance mensuelle."),
        ("Ranking", "Which Fekola machines have the lowest Availability?", "Classement equipements."),
        ("Downtime", "Show the top downtime drivers for HMS at Fekola YTD.", "Drivers + Pareto sur product_group HMS."),
        ("Work Type", "Show unplanned downtime drivers at Fekola.", "Filtre WorkType=Unplanned."),
        ("Parts", "YTD Parts Sales for Fekola.", "CA Facture EU, LOB PARTS, scope customer/territory."),
        ("Knowledge", "What is MTBF?", "Definition validee, sans requete KPI operationnelle."),
        ("Combined", "Why is Fekola MTBF low and what should we do?", "Constats + evidence + Best Practices."),
        ("Capability", "Que peux-tu faire pour moi ?", "Catalogue personnalise sans Power BI/LLM."),
        ("Answerability", "Quelle est la date de mise en service de DT677 ?", "Abstention champ non configure."),
        ("Follow-up", "Only the 777.", "Herite le site/periode et remplace le modele."),
        ("Export", "Download it.", "Exporte le dernier artefact compatible."),
        ("Reporting", "Open the Fleet Performance report.", "Navigation vers le rapport autorise."),
    ]
    add_table(doc, ["Famille", "Question", "Resultat attendu"], examples, widths=[3, 7, 7], font_size=7)

    add_heading(doc, "Annexe G. Glossaire")
    add_table(doc, ["Terme", "Definition"], [
        ("Intent", "Representation structuree de l'objectif analytique."),
        ("Entity", "Valeur metier resolue : site, modele, equipement, serial, KPI, periode."),
        ("Capability", "Fonction gouvernee exposee selon readiness et permissions."),
        ("Semantic model", "Modele Power BI contenant mesures, dimensions et RLS."),
        ("DAX template", "Patron ferme utilise pour generer une requete controlee."),
        ("Artifact", "Snapshot persistant et immutable d'un resultat conversationnel."),
        ("Answerability", "Decision backend indiquant si une reponse fiable est possible."),
        ("Grounding", "Verification que les faits affiches existent dans l'evidence."),
        ("RLS", "Row-Level Security Power BI limitant les lignes accessibles."),
        ("Effective identity", "Identite transmise au modele pour appliquer le RLS."),
        ("YTD", "Year to Date, du debut de l'annee a la date de donnees disponible."),
        ("Physical Availability", "Mesure [Availability New], affichee en pourcentage."),
        ("MTBS", "Mean Time Between Stoppages."),
        ("MTBF", "Mean Time Between Failures."),
        ("MTTR", "Mean Time To Repair."),
        ("SMU", "Service Meter Unit issue de l'equipement master; unite a confirmer."),
        ("Pareto", "Classement des drivers avec contribution et cumul."),
        ("Feature flag", "Parametre de rollout Disabled/Admin/Pilot/Production."),
    ], widths=[4, 12], font_size=7.5)


def finalize(doc):
    doc.core_properties.title = "Chatbot Mining 360 - Reference complete"
    doc.core_properties.subject = "Capacites, parametres, architecture, securite et exploitation"
    doc.core_properties.author = "Mining360"
    doc.core_properties.keywords = "Mining360, chatbot, Power BI, Fleet, Availability, RLS, AI"
    doc.core_properties.comments = "Generated from the Mining360IA implementation and active configuration."
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


def main():
    doc = Document()
    configure_document(doc)
    add_title(doc)
    document_control(doc)
    executive_summary(doc)
    architecture(doc)
    user_experience(doc)
    agents_and_routing(doc)
    intent_and_parameters(doc)
    entity_and_time(doc)
    fleet_inventory(doc)
    fleet_performance(doc)
    availability(doc)
    downtime(doc)
    parts_sales(doc)
    knowledge_and_reporting(doc)
    capability_answerability(doc)
    persistence_exports(doc)
    security(doc)
    config_admin(doc)
    production_hardening(doc)
    feature_flag_section(doc)
    api_contracts(doc)
    errors_observability(doc)
    performance_testing(doc)
    limitations(doc)
    appendices(doc)
    finalize(doc)
    print(OUTPUT)


if __name__ == "__main__":
    main()
